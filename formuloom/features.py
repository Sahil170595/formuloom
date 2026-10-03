from __future__ import annotations

import logging
import re
import time as _time
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Literal

from openpyxl.formula.tokenizer import Tokenizer  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]
from openpyxl.utils.cell import SHEETRANGE_RE, range_boundaries  # type: ignore[import-untyped]

from formuloom.bundle import TaskBundle
from formuloom.schema import CellEntry, cell_column_index, cell_row
from formuloom.workbook import CellRecord, SheetData, WorkbookData

logger = logging.getLogger(__name__)

SECTION_BLANK_ROWS_TERMINATE = 2

DEFINED_NAME_MAX_DEPTH = 5

AGG_FUNCTIONS: frozenset[str] = frozenset({"SUM", "SUBTOTAL", "AGGREGATE", "AVERAGE", "SUMIF"})

DOMAIN_LEXICON: tuple[str, ...] = (
    "total",
    "subtotal",
    "net",
    "gross",
    "ebitda",
    "ebit",
    "net income",
    "net profit",
    "eps",
    "enterprise value",
    "equity value",
    "implied share price",
    "share price",
    "price per share",
    "irr",
    "moic",
    "noi",
    "dscr",
    "ending cash",
    "closing cash",
    "cash balance",
    "accretion",
    "dilution",
    "free cash flow",
    "terminal value",
    "wacc",
    "npv",
    "net present value",
    "gross profit",
    "gross margin",
    "operating income",
    "operating profit",
    "net revenue",
    "net debt",
    "total assets",
    "total liabilities",
    "total equity",
    "valuation",
    "exit value",
    "payback",
    "ltv",
    "cap rate",
    "debt service",
    "check",
    "balance",
)

_LEXICON_PATTERNS: tuple[tuple[str, re.Pattern[str]], ...] = tuple(
    (term, re.compile(rf"\b{re.escape(term)}\b", re.IGNORECASE)) for term in DOMAIN_LEXICON
)

CellKey = tuple[str, str]

PatternKind = Literal["r1c1", "sketch", "mixed", "none"]

_A1_CELL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)([0-9]+)$")
_A1_COL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})$")
_A1_ROW_RE = re.compile(r"^(\$?)([0-9]+)$")
_NAME_RE = re.compile(r"^[A-Za-z_\\][A-Za-z0-9_.\\]*$")
_XLFN_PREFIX = "_XLFN."
_LAMBDA_PARAM_PREFIX = "_xlpm."


class FeatureError(Exception):
    pass


class SheetNameMismatchError(FeatureError):
    pass


@dataclass(frozen=True)
class Section:

    sheet: str
    title: str | None
    heading_row: int | None
    start_row: int
    end_row: int


@dataclass
class ReferenceGraph:

    precedents: dict[CellKey, set[CellKey]]
    dependents: dict[CellKey, set[CellKey]]
    warnings: dict[CellKey, list[str]]


@dataclass
class CellFeatures:

    sheet: str
    ref: str
    row: int
    col: int
    cell_type: str
    label: str | None
    label_ref: str | None
    value: object | None
    formula: str | None
    sketch: str | None
    r1c1: str | None
    is_input_colored: bool
    bold_or_bordered: bool
    n_dependents: int
    warnings: list[str] = field(default_factory=list)


@dataclass
class RowFeatures:

    sheet: str
    row: int
    label: str | None
    section_title: str | None
    n_candidates: int
    candidate_refs: list[str]
    n_dependents_outside_row: int
    distinct_source_sheets: int
    distinct_source_sections: int
    aggregates_range: bool
    agg_function: str | None
    lexicon_hits: list[str]
    input_colored_fraction: float
    bold_or_bordered: bool
    formula_pattern: str | None
    pattern_kind: PatternKind
    in_chart_source: bool
    in_print_area: bool
    in_named_range: bool
    n_distinct_labels_left: int
    n_distinct_sketches: int
    warnings: list[str] = field(default_factory=list)


@dataclass
class SheetFeatures:

    sheet: str
    sections: list[Section]
    rows: list[RowFeatures]
    cells: dict[str, CellFeatures]
    n_candidates: int
    used_range: str
    max_row: int
    max_col: int
    candidate_density: float
    frac_rows_multi_label: float
    frac_rows_multi_sketch: float
    label_value_alternation: float
    input_colored_fraction: float
    zero_formula: bool
    n_cells_with_cross_sheet_precedents: int
    n_cells_with_cross_sheet_dependents: int


@dataclass
class TaskFeatures:

    task: str
    graph: ReferenceGraph
    sheets: dict[str, SheetFeatures]


def _split_sheet_prefix(operand: str) -> tuple[str | None, str]:
    match = SHEETRANGE_RE.match(operand)
    if match is None or match.end() != len(operand):
        return None, operand
    quoted = match.group("quoted")
    sheet = quoted.replace("''", "'") if quoted is not None else match.group("notquoted")
    return sheet, match.group("cells")


def _split_top_level_commas(target: str) -> list[str]:
    parts: list[str] = []
    buf: list[str] = []
    in_quote = False
    for char in target:
        if char == "'":
            in_quote = not in_quote
        if char == "," and not in_quote:
            parts.append("".join(buf))
            buf = []
        else:
            buf.append(char)
    parts.append("".join(buf))
    return [p.strip() for p in parts if p.strip()]


def _expand_bounds(
    sheet_name: str,
    text: str,
    bounds: tuple[int | None, int | None, int | None, int | None],
    workbook: WorkbookData,
    warnings: list[str],
) -> set[CellKey]:
    target = workbook.sheets.get(sheet_name)
    if target is None:
        warnings.append(f"ref {text!r} names unknown sheet {sheet_name!r}")
        return set()
    min_col, min_row, max_col, max_row = bounds
    capped = min_row is None or max_row is None or min_col is None or max_col is None
    min_col = 1 if min_col is None else min_col
    min_row = 1 if min_row is None else min_row
    max_col = target.max_col if max_col is None else min(max_col, target.max_col)
    max_row = target.max_row if max_row is None else min(max_row, target.max_row)
    if capped:
        warnings.append(f"whole-column/row ref {text!r} capped to used range of {sheet_name!r}")
    if max_row < min_row or max_col < min_col:
        return set()
    return {
        (sheet_name, f"{get_column_letter(col)}{row}")
        for row in range(min_row, max_row + 1)
        for col in range(min_col, max_col + 1)
    }


def _resolve_defined_name(
    name: str, current_sheet: str, workbook: WorkbookData, warnings: list[str], depth: int
) -> set[CellKey]:
    target = workbook.defined_names.get(name)
    if target is None:
        folded = name.casefold()
        target = next((t for n, t in workbook.defined_names.items() if n.casefold() == folded), None)
    if target is None:
        warnings.append(f"unresolvable ref {name!r} (not a cell, range, or defined name)")
        return set()
    if depth >= DEFINED_NAME_MAX_DEPTH:
        warnings.append(f"defined name {name!r} exceeds resolution depth {DEFINED_NAME_MAX_DEPTH}")
        return set()
    out: set[CellKey] = set()
    for area in _split_top_level_commas(target):
        out |= _resolve_operand(area, current_sheet, workbook, warnings, depth + 1)
    return out


def _resolve_operand(
    operand: str, current_sheet: str, workbook: WorkbookData, warnings: list[str], depth: int = 0
) -> set[CellKey]:
    op = operand.strip()
    if op.startswith(_LAMBDA_PARAM_PREFIX):
        return set()
    if not op or "#REF!" in op or op.startswith("#"):
        warnings.append(f"unresolvable ref {op!r} (error ref)")
        return set()
    sheet, cells = _split_sheet_prefix(op)
    if sheet is not None:
        if "[" in sheet:
            warnings.append(f"external workbook ref {op!r} not resolvable")
            return set()
        try:
            bounds = range_boundaries(cells)
        except ValueError:
            warnings.append(f"unparseable ref {op!r}")
            return set()
        return _expand_bounds(sheet, op, bounds, workbook, warnings)
    if "[" in op:
        warnings.append(f"structured ref {op!r} not resolvable")
        return set()

    if _A1_CELL_RE.match(op) or ":" in op:
        try:
            bounds = range_boundaries(op)
        except ValueError:
            warnings.append(f"unparseable ref {op!r}")
            return set()
        return _expand_bounds(current_sheet, op, bounds, workbook, warnings)
    return _resolve_defined_name(op, current_sheet, workbook, warnings, depth)


def _tokenize(formula: str, warnings: list[str]) -> list[object] | None:
    try:
        return list(Tokenizer(formula).items)
    except Exception as exc:  # noqa: BLE001 - tokenizer raises mixed exception types on bad input
        logger.debug("formula tokenization failed for %r: %s", formula, exc)
        warnings.append(f"formula tokenization failed: {exc}")
        return None


def _extract_refs(formula: str, current_sheet: str, workbook: WorkbookData, warnings: list[str]) -> set[CellKey]:
    tokens = _tokenize(formula, warnings)
    if tokens is None:
        return set()
    out: set[CellKey] = set()
    for tok in tokens:
        if getattr(tok, "type", None) == "OPERAND" and getattr(tok, "subtype", None) == "RANGE":
            out |= _resolve_operand(str(getattr(tok, "value", "")), current_sheet, workbook, warnings)
    return out


def build_reference_graph(workbook: WorkbookData) -> ReferenceGraph:
    started = _time.perf_counter()
    precedents: dict[CellKey, set[CellKey]] = {}
    dependents: dict[CellKey, set[CellKey]] = {}
    warnings: dict[CellKey, list[str]] = {}
    n_formulas = 0
    for sheet_name, sheet in workbook.sheets.items():
        for ref, rec in sheet.cells.items():
            if rec.formula is None:
                continue
            n_formulas += 1
            key = (sheet_name, ref)
            cell_warnings: list[str] = []
            reads = _extract_refs(rec.formula, sheet_name, workbook, cell_warnings)
            precedents[key] = reads
            for target in reads:
                dependents.setdefault(target, set()).add(key)
            if cell_warnings:
                warnings[key] = cell_warnings
    logger.info(
        "reference graph for %s: %d formulas, %d cells warned, %.2fs",
        workbook.path.name,
        n_formulas,
        len(warnings),
        _time.perf_counter() - started,
    )
    return ReferenceGraph(precedents=precedents, dependents=dependents, warnings=warnings)


def _rel_part(prefix: str, offset: int) -> str:
    return prefix if offset == 0 else f"{prefix}[{offset}]"


def _corner_to_r1c1(part: str, row: int, col: int) -> str | None:
    cell = _A1_CELL_RE.match(part)
    if cell is not None:
        col_abs, letters, row_abs, digits = cell.groups()
        col_idx = _col_index(letters)
        r_part = f"R{digits}" if row_abs else _rel_part("R", int(digits) - row)
        c_part = f"C{col_idx}" if col_abs else _rel_part("C", col_idx - col)
        return r_part + c_part
    col_only = _A1_COL_RE.match(part)
    if col_only is not None:
        col_abs, letters = col_only.groups()
        col_idx = _col_index(letters)
        return f"C{col_idx}" if col_abs else _rel_part("C", col_idx - col)
    row_only = _A1_ROW_RE.match(part)
    if row_only is not None:
        row_abs, digits = row_only.groups()
        return f"R{digits}" if row_abs else _rel_part("R", int(digits) - row)
    return None


def _col_index(letters: str) -> int:
    idx = 0
    for char in letters.upper():
        idx = idx * 26 + (ord(char) - ord("A") + 1)
    return idx


def _range_operand_to_r1c1(operand: str, row: int, col: int) -> str | None:
    if "[" in operand:
        return None
    sheet, cells = _split_sheet_prefix(operand)
    prefix = operand[: len(operand) - len(cells)] if sheet is not None else ""
    body = cells if sheet is not None else operand
    if sheet is None and ":" not in body and not _A1_CELL_RE.match(body):

        return operand if _NAME_RE.match(body) else None
    corners = body.split(":")
    if len(corners) > 2 or not body:
        return None
    if len(corners) == 1 and not _A1_CELL_RE.match(body):
        return None
    converted: list[str] = []
    for corner in corners:
        r1c1 = _corner_to_r1c1(corner, row, col)
        if r1c1 is None:
            return None
        converted.append(r1c1)
    if len(converted) == 2 and converted[0] == converted[1]:
        converted = converted[:1]
    return prefix + ":".join(converted)


def to_r1c1(formula: str, row: int, col: int) -> str | None:
    body = formula[1:] if formula.startswith("=") else formula
    if not body:
        return None
    try:
        tokens = list(Tokenizer("=" + body).items)
    except Exception as exc:  # noqa: BLE001 - tokenizer raises mixed exception types on bad input
        logger.debug("to_r1c1 tokenization failed for %r: %s", formula, exc)
        return None
    parts: list[str] = []
    for tok in tokens:
        value = str(tok.value)
        if tok.type == "OPERAND" and tok.subtype == "RANGE":
            converted = _range_operand_to_r1c1(value, row, col)
            if converted is None:
                return None
            parts.append(converted)
        else:
            parts.append(value)
    return "=" + "".join(parts)


def formula_sketch(formula: str) -> str:
    body = formula[1:] if formula.startswith("=") else formula
    try:
        tokens = list(Tokenizer("=" + body).items) if body else []
    except Exception as exc:  # noqa: BLE001 - fallback path IS the contract (always succeed)
        logger.debug("sketch tokenization failed for %r; using regex fallback: %s", formula, exc)
        masked = re.sub(r'"[^"]*"', "STR", body)
        masked = re.sub(r"\$?[A-Za-z]{1,3}\$?[0-9]+", "REF", masked)
        masked = re.sub(r"\d+(?:\.\d+)?(?:[Ee][+-]?\d+)?", "NUM", masked)
        return re.sub(r"\s+", "", masked)
    parts: list[str] = []
    for tok in tokens:
        value = str(tok.value)
        if tok.type == "OPERAND" and tok.subtype == "RANGE":
            parts.append("REF:REF" if ":" in value else "REF")
        elif tok.type == "OPERAND" and tok.subtype == "NUMBER":
            parts.append("NUM")
        elif tok.type == "OPERAND" and tok.subtype == "TEXT":
            parts.append("STR")
        elif tok.type == "WHITE-SPACE":
            parts.append(" ")
        else:
            parts.append(value)
    return "".join(parts).strip()


def _top_level_aggregate(formula: str) -> str | None:
    try:
        tokens = [t for t in Tokenizer(formula).items if t.type != "WHITE-SPACE"]
    except Exception as exc:  # noqa: BLE001 - aggregation shape is a best-effort feature
        logger.debug("aggregate detection failed for %r: %s", formula, exc)
        return None
    if not tokens or tokens[0].type != "FUNC" or tokens[0].subtype != "OPEN":
        return None
    fname = str(tokens[0].value)[:-1].upper()
    fname = fname.removeprefix(_XLFN_PREFIX)
    if fname not in AGG_FUNCTIONS:
        return None
    depth = 1
    saw_range = False
    for idx, tok in enumerate(tokens[1:], start=1):
        if tok.subtype == "OPEN":
            depth += 1
        elif tok.subtype == "CLOSE":
            depth -= 1
            if depth == 0 and idx != len(tokens) - 1:
                return None
        elif depth == 1 and tok.type == "OPERAND" and tok.subtype == "RANGE" and ":" in str(tok.value):
            saw_range = True
    return fname if saw_range else None


def _rows_index(sheet: SheetData) -> dict[int, list[CellRecord]]:
    out: dict[int, list[CellRecord]] = {}
    for rec in sheet.cells.values():
        out.setdefault(rec.row, []).append(rec)
    for cells in out.values():
        cells.sort(key=lambda c: c.col)
    return out


def _is_numeric_value(value: object) -> bool:
    if isinstance(value, bool):
        return False
    return isinstance(value, (int, float, datetime, date, time))


def _row_text_of(cells: list[CellRecord]) -> str | None:
    for rec in cells:
        if isinstance(rec.value, str) and rec.value.strip():
            return rec.value.strip()
    return None


def _is_value_cell(rec: CellRecord) -> bool:
    if _is_numeric_value(rec.value):
        return True
    return rec.formula is not None and not isinstance(rec.value, str)


def _is_heading_row(cells: list[CellRecord], next_cells: list[CellRecord]) -> bool:
    if any(_is_value_cell(rec) for rec in cells):
        return False
    leftmost = cells[0]
    if not (isinstance(leftmost.value, str) and leftmost.value.strip()):
        return False
    text_cols = [rec.col for rec in cells if isinstance(rec.value, str) and rec.value.strip()]
    if text_cols != list(range(text_cols[0], text_cols[0] + len(text_cols))):
        return False
    title = leftmost.value.strip()
    bold = any(rec.bold for rec in cells)
    all_caps = title.isupper() and any(ch.isalpha() for ch in title)
    followed_by_blank = not next_cells
    underline = any(rec.border_bottom for rec in cells) or any(rec.border_top for rec in next_cells)
    return bold or all_caps or followed_by_blank or underline


def detect_sections(sheet: SheetData) -> list[Section]:
    rows = _rows_index(sheet)
    if not rows:
        return []
    sections: list[Section] = []
    start: int | None = None
    heading: int | None = None
    title: str | None = None
    last_nonblank = 0
    blanks = 0

    def close() -> None:
        nonlocal start, heading, title
        if start is not None:
            sections.append(
                Section(sheet=sheet.name, title=title, heading_row=heading, start_row=start, end_row=last_nonblank)
            )
        start, heading, title = None, None, None

    for row_num in range(1, sheet.max_row + 1):
        cells = rows.get(row_num)
        if not cells:
            blanks += 1
            if blanks >= SECTION_BLANK_ROWS_TERMINATE:
                close()
            continue
        blanks = 0
        if _is_heading_row(cells, rows.get(row_num + 1, [])):
            close()
            start, heading, title = row_num, row_num, _row_text_of(cells)
        elif start is None:
            start, heading, title = row_num, None, None
        last_nonblank = row_num
    close()
    return sections


def _section_map(sections: list[Section]) -> dict[int, int]:
    out: dict[int, int] = {}
    for idx, section in enumerate(sections):
        for row_num in range(section.start_row, section.end_row + 1):
            out[row_num] = idx
    return out


Bounds = tuple[int, int, int, int]


def _areas_from_target(target: str, default_sheet: str | None, workbook: WorkbookData) -> list[tuple[str, Bounds]]:
    areas: list[tuple[str, Bounds]] = []
    for part in _split_top_level_commas(target):
        sheet, cells = _split_sheet_prefix(part)
        if sheet is None:
            sheet, cells = default_sheet, part
        if sheet is None or "[" in sheet or "#REF!" in cells:
            continue
        try:
            min_col, min_row, max_col, max_row = range_boundaries(cells)
        except ValueError:
            logger.debug("skipping unparseable presentation area %r", part)
            continue
        target_sheet = workbook.sheets.get(sheet)
        if target_sheet is None:
            continue
        areas.append(
            (
                sheet,
                (
                    min_col or 1,
                    min_row or 1,
                    max_col or target_sheet.max_col,
                    max_row or target_sheet.max_row,
                ),
            )
        )
    return areas


@dataclass
class _PresentationAreas:
    chart: list[tuple[str, Bounds]]
    print_area: list[tuple[str, Bounds]]
    named: list[tuple[str, Bounds]]

    @classmethod
    def collect(cls, workbook: WorkbookData) -> _PresentationAreas:
        chart: list[tuple[str, Bounds]] = []
        for refs in workbook.chart_source_refs.values():
            for ref in refs:
                chart.extend(_areas_from_target(ref, None, workbook))
        print_area: list[tuple[str, Bounds]] = []
        for sheet_name, area in workbook.print_areas.items():
            if area:
                print_area.extend(_areas_from_target(area, sheet_name, workbook))
        named: list[tuple[str, Bounds]] = []
        for target in workbook.defined_names.values():
            named.extend(_areas_from_target(target, None, workbook))
        return cls(chart=chart, print_area=print_area, named=named)


def _in_areas(areas: list[tuple[str, Bounds]], sheet: str, row: int, col: int) -> bool:
    return any(s == sheet and b[0] <= col <= b[2] and b[1] <= row <= b[3] for s, b in areas)


def _nearest_left_label(
    row_cells: list[CellRecord], col: int, candidate_cols: set[int]
) -> tuple[str | None, str | None]:
    best: CellRecord | None = None
    best_candidate: CellRecord | None = None
    for rec in row_cells:
        if rec.col >= col:
            break
        if not (isinstance(rec.value, str) and rec.value.strip() and any(ch.isalpha() for ch in rec.value)):
            continue
        if rec.col in candidate_cols:
            best_candidate = rec
        else:
            best = rec
    chosen = best if best is not None else best_candidate
    if chosen is None:
        return None, None
    return chosen.value.strip() if isinstance(chosen.value, str) else None, chosen.ref


def _lexicon_hits(label: str | None) -> list[str]:
    if not label:
        return []
    return [term for term, pattern in _LEXICON_PATTERNS if pattern.search(label)]


def _pattern_for(cands: list[CellFeatures]) -> tuple[PatternKind, str | None]:
    with_formula = [c for c in cands if c.formula is not None]
    if not with_formula:
        return "none", None
    r1c1s = {c.r1c1 for c in with_formula}
    if None not in r1c1s and len(r1c1s) == 1:
        only = next(iter(r1c1s))
        return "r1c1", only
    sketches = {c.sketch for c in with_formula}
    if len(sketches) == 1:
        return "sketch", next(iter(sketches))
    return "mixed", None


def _near_miss_suggestions(wanted: str, available: list[str]) -> list[str]:
    folded = wanted.strip().casefold()
    return [name for name in available if name.strip().casefold() == folded]


def _check_sheet_names(candidates: dict[str, list[CellEntry]], workbook: WorkbookData) -> None:
    available = list(workbook.sheets)
    for sheet_name in candidates:
        if sheet_name not in workbook.sheets:
            near = _near_miss_suggestions(sheet_name, available)
            hint = f" near-miss candidates: {near}" if near else ""
            raise SheetNameMismatchError(
                f"raw_diff sheet {sheet_name!r} not found in workbook "
                f"{workbook.path.name} (sheets: {available}).{hint}"
            )


def _cross_sheet_cell_counts(graph: ReferenceGraph) -> dict[str, tuple[int, int]]:
    reading: dict[str, set[CellKey]] = {}
    read_from: dict[str, set[CellKey]] = {}
    for key, reads in graph.precedents.items():
        sheet = key[0]
        for target in reads:
            if target[0] != sheet:
                reading.setdefault(sheet, set()).add(key)
                read_from.setdefault(target[0], set()).add(target)
    return {s: (len(reading.get(s, ())), len(read_from.get(s, ()))) for s in set(reading) | set(read_from)}


def _build_cell_features(
    entry: CellEntry,
    sheet: SheetData,
    row_cells: list[CellRecord],
    candidate_cols: set[int],
    graph: ReferenceGraph,
) -> CellFeatures:
    ref = entry.cell
    row = cell_row(ref)
    col = cell_column_index(ref)
    rec = sheet.cells.get(ref)
    label, label_ref = _nearest_left_label(row_cells, col, candidate_cols)
    formula = rec.formula if rec is not None else None
    warnings = list(graph.warnings.get((sheet.name, ref), []))
    return CellFeatures(
        sheet=sheet.name,
        ref=ref,
        row=row,
        col=col,
        cell_type=entry.cell_type,
        label=label,
        label_ref=label_ref,
        value=rec.value if rec is not None else None,
        formula=formula,
        sketch=formula_sketch(formula) if formula is not None else None,
        r1c1=to_r1c1(formula, row, col) if formula is not None else None,
        is_input_colored=rec.is_input_colored if rec is not None else False,
        bold_or_bordered=(rec.bold or rec.border_top or rec.border_bottom) if rec is not None else False,
        n_dependents=len(graph.dependents.get((sheet.name, ref), ())),
        warnings=warnings,
    )


def _build_row_features(
    sheet: SheetData,
    row_num: int,
    cands: list[CellFeatures],
    row_cells: list[CellRecord],
    graph: ReferenceGraph,
    sections: list[Section],
    section_maps: dict[str, dict[int, int]],
    areas: _PresentationAreas,
) -> RowFeatures:
    section_idx = section_maps.get(sheet.name, {}).get(row_num)
    section_title = sections[section_idx].title if section_idx is not None else None

    outside_dependents: set[CellKey] = set()
    source_sheets: set[str] = set()
    source_sections: set[tuple[str, int]] = set()
    agg_function: str | None = None
    warnings: list[str] = []
    for cand in cands:
        key = (sheet.name, cand.ref)
        for dep_sheet, dep_ref in graph.dependents.get(key, ()):
            if not (dep_sheet == sheet.name and cell_row(dep_ref) == row_num):
                outside_dependents.add((dep_sheet, dep_ref))
        for src_sheet, src_ref in graph.precedents.get(key, ()):
            source_sheets.add(src_sheet)
            src_section = section_maps.get(src_sheet, {}).get(cell_row(src_ref), -1)
            source_sections.add((src_sheet, src_section))
        if agg_function is None and cand.formula is not None:
            agg_function = _top_level_aggregate(cand.formula)
        warnings.extend(cand.warnings)

    candidate_cols = {c.col for c in cands}
    label, _label_ref = _nearest_left_label(row_cells, cands[0].col, candidate_cols)
    pattern_kind, formula_pattern = _pattern_for(cands)
    stored = [c for c in cands if (sheet.cells.get(c.ref)) is not None]
    colored = sum(1 for c in stored if c.is_input_colored)
    return RowFeatures(
        sheet=sheet.name,
        row=row_num,
        label=label,
        section_title=section_title,
        n_candidates=len(cands),
        candidate_refs=[c.ref for c in cands],
        n_dependents_outside_row=len(outside_dependents),
        distinct_source_sheets=len(source_sheets),
        distinct_source_sections=len(source_sections),
        aggregates_range=agg_function is not None,
        agg_function=agg_function,
        lexicon_hits=_lexicon_hits(label),
        input_colored_fraction=colored / len(stored) if stored else 0.0,
        bold_or_bordered=any(c.bold_or_bordered for c in cands),
        formula_pattern=formula_pattern,
        pattern_kind=pattern_kind,
        in_chart_source=any(_in_areas(areas.chart, sheet.name, c.row, c.col) for c in cands),
        in_print_area=any(_in_areas(areas.print_area, sheet.name, c.row, c.col) for c in cands),
        in_named_range=any(_in_areas(areas.named, sheet.name, c.row, c.col) for c in cands),
        n_distinct_labels_left=len({c.label_ref for c in cands if c.label_ref is not None}),
        n_distinct_sketches=len({c.sketch for c in cands if c.sketch is not None}),
        warnings=warnings,
    )


def _label_value_alternation(
    rows_index: dict[int, list[CellRecord]], candidate_cols_by_row: dict[int, set[int]]
) -> float:
    text_cols: dict[int, int] = {}
    value_cols: dict[int, int] = {}
    for row_num, candidate_cols in candidate_cols_by_row.items():
        for rec in rows_index.get(row_num, []):
            is_candidate = rec.col in candidate_cols
            is_label_text = isinstance(rec.value, str) and rec.value.strip() and any(ch.isalpha() for ch in rec.value)
            if is_label_text and not is_candidate:
                text_cols[rec.col] = text_cols.get(rec.col, 0) + 1
            elif is_candidate or _is_numeric_value(rec.value):
                value_cols[rec.col] = value_cols.get(rec.col, 0) + 1
    transitions = 0
    previous: str | None = None
    for col in sorted(set(text_cols) | set(value_cols)):
        kind = "text" if text_cols.get(col, 0) >= value_cols.get(col, 0) else "value"
        if previous == "text" and kind == "value":
            transitions += 1
        previous = kind
    return float(transitions)


def _build_sheet_features(
    sheet: SheetData,
    entries: list[CellEntry],
    graph: ReferenceGraph,
    sections_by_sheet: dict[str, list[Section]],
    section_maps: dict[str, dict[int, int]],
    areas: _PresentationAreas,
    cross_edges: dict[str, tuple[int, int]],
) -> SheetFeatures:
    rows_index = _rows_index(sheet)
    sections = sections_by_sheet.get(sheet.name, [])

    ordered = sorted(entries, key=lambda e: (cell_row(e.cell), cell_column_index(e.cell)))
    candidate_cols_by_row: dict[int, set[int]] = {}
    for entry in ordered:
        candidate_cols_by_row.setdefault(cell_row(entry.cell), set()).add(cell_column_index(entry.cell))

    cells: dict[str, CellFeatures] = {}
    by_row: dict[int, list[CellFeatures]] = {}
    for entry in ordered:
        row_num = cell_row(entry.cell)
        feats = _build_cell_features(entry, sheet, rows_index.get(row_num, []), candidate_cols_by_row[row_num], graph)
        cells[entry.cell] = feats
        by_row.setdefault(feats.row, []).append(feats)

    rows = [
        _build_row_features(sheet, row_num, cands, rows_index.get(row_num, []), graph, sections, section_maps, areas)
        for row_num, cands in sorted(by_row.items())
    ]

    n_candidates = len(cells)
    used_cells = max(1, sheet.max_row * sheet.max_col)
    stored = [c for c in cells.values() if sheet.cells.get(c.ref) is not None]
    colored = sum(1 for c in stored if c.is_input_colored)
    n_rows = max(1, len(rows))
    out_edges, in_edges = cross_edges.get(sheet.name, (0, 0))
    return SheetFeatures(
        sheet=sheet.name,
        sections=sections,
        rows=rows,
        cells=cells,
        n_candidates=n_candidates,
        used_range=f"A1:{get_column_letter(max(1, sheet.max_col))}{max(1, sheet.max_row)}",
        max_row=sheet.max_row,
        max_col=sheet.max_col,
        candidate_density=n_candidates / used_cells,
        frac_rows_multi_label=sum(1 for r in rows if r.n_distinct_labels_left > 1) / n_rows,
        frac_rows_multi_sketch=sum(1 for r in rows if r.n_distinct_sketches > 1) / n_rows,
        label_value_alternation=_label_value_alternation(rows_index, candidate_cols_by_row),
        input_colored_fraction=colored / len(stored) if stored else 0.0,
        zero_formula=all(c.formula is None for c in cells.values()),
        n_cells_with_cross_sheet_precedents=out_edges,
        n_cells_with_cross_sheet_dependents=in_edges,
    )


def build_task_features(bundle: TaskBundle, complete: WorkbookData) -> TaskFeatures:
    candidates = bundle.candidate_cells()
    _check_sheet_names(candidates, complete)
    graph = build_reference_graph(complete)
    sections_by_sheet = {name: detect_sections(sheet) for name, sheet in complete.sheets.items()}
    section_maps = {name: _section_map(sections) for name, sections in sections_by_sheet.items()}
    areas = _PresentationAreas.collect(complete)
    cross_edges = _cross_sheet_cell_counts(graph)
    sheets = {
        name: _build_sheet_features(
            sheet, candidates.get(name, []), graph, sections_by_sheet, section_maps, areas, cross_edges
        )
        for name, sheet in complete.sheets.items()
    }
    return TaskFeatures(task=bundle.name, graph=graph, sheets=sheets)
