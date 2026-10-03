"""Fresh synthetic XLSX bundles; numeric formula caches are authored, not recalculated."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]
from openpyxl.styles import Font  # type: ignore[import-untyped]

from formuloom.assemble import assemble_diff_file
from formuloom.perturb import _patch_cached_values
from formuloom.schema import CellEntry, CellGroup, DiffFile, SheetDiff, SheetGroups

FIXTURE_MARKER = "Synthetic public fixture: operations-rollup-v1"
FIXTURE_TIMESTAMP = datetime(2026, 1, 1)


def generate_fixture(root: Path) -> Path:
    """Create a new bundle without overwriting an existing directory."""
    root = Path(root)
    root.mkdir(parents=True, exist_ok=False)
    candidates = {
        "Assumptions": ["B2", "B3"],
        "Operations": ["B2", "C2", "B4", "C4", "B5", "C5", "B6", "C6", "B7", "C7"],
        "Summary": ["B2", "C2", "B3", "C3", "B4", "C4"],
    }
    for name, units in (("init.xlsx", 200), ("complete.xlsx", 220)):
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        wb.properties.creator = "Formuloom"
        wb.properties.lastModifiedBy = "Formuloom"
        wb.properties.created = FIXTURE_TIMESTAMP
        wb.properties.modified = FIXTURE_TIMESTAMP
        wb.properties.description = "Independently generated synthetic public workbook."
        assumptions = wb.create_sheet("Assumptions")
        assumptions.append(["OPERATING INPUTS"])
        assumptions.append(["Unit price", 3 if units == 220 else 2])
        assumptions.append(["Fixed cost", 100 if units == 220 else 90])
        assumptions["B2"].font = Font(color="0000FF")
        operations = wb.create_sheet("Operations")
        operations.append(["PERIOD OPERATIONS", "Period A", "Period B"])
        operations.append(["Units delivered", units, units + 10])
        operations.append(["Price per unit", "=Assumptions!B2", "=Assumptions!B2"])
        operations.append(["Gross proceeds", "=B2*B3", "=C2*C3"])
        operations.append(["Total revenue", "=SUM(B4:B4)", "=SUM(C4:C4)"])
        operations.append(["Fixed cost", "=Assumptions!B3", "=Assumptions!B3"])
        operations.append(["Net contribution", "=B5-B6", "=C5-C6"])
        summary = wb.create_sheet("Summary")
        summary.append(["READER RESULTS", "Period A", "Period B"])
        summary.append(["Total revenue", "=Operations!B5", "=Operations!C5"])
        summary.append(["Net contribution", "=Operations!B7", "=Operations!C7"])
        summary.append(["Total reviewed results", "=SUM(B2:B3)", "=SUM(C2:C3)"])
        for ws in wb:
            ws.freeze_panes = "B2"
            ws.column_dimensions["A"].width = 28
            for row in ws.iter_rows():
                for cell in row:
                    if cell.row == 1 or cell.row in (5, 7) or ws.title == "Summary":
                        cell.font = Font(bold=True)
                    if cell.column > 1 and cell.row > 1:
                        cell.number_format = "0.00"
        price, cost = (3, 100) if units == 220 else (2, 90)
        gross = [units * price, (units + 10) * price]
        cached: dict[str, dict[str, object]] = {"Operations": {}, "Summary": {}}
        for col, amount in zip(("B", "C"), gross, strict=True):
            cached["Operations"].update(
                {f"{col}3": price, f"{col}4": amount, f"{col}5": amount, f"{col}6": cost, f"{col}7": amount - cost}
            )
            cached["Summary"].update({f"{col}2": amount, f"{col}3": amount - cost, f"{col}4": 2 * amount - cost})
        path = root / name
        wb.save(path)
        _patch_cached_values(path, cached, epoch=wb.epoch)
        wb.close()
    raw = DiffFile(
        sheets={
            sheet: SheetDiff(
                sheet_weight=1,
                groups=SheetGroups(
                    intermediate=CellGroup(
                        weight=1,
                        cells=[CellEntry(cell=ref, cell_type="currency") for ref in refs],
                    )
                ),
            )
            for sheet, refs in candidates.items()
        }
    )
    reference = assemble_diff_file(
        raw, {"Assumptions": set(), "Operations": {"B5", "C5", "B7", "C7"}, "Summary": set(candidates["Summary"])}
    )
    for filename, payload in (("raw_diff.json", raw), ("subset.json", reference)):
        (root / filename).write_text(json.dumps(payload.model_dump(), indent=2) + "\n", encoding="utf-8")
    (root / "instructions.md").write_text(
        FIXTURE_MARKER + "\n\nReview total revenue and net contribution for each period, "
        "plus every row of the reader-results sheet. Inputs and gross-proceeds precursor "
        "calculations are intermediate. This tiny policy illustrates mechanics, not generalization.\n",
        encoding="utf-8",
    )
    return root
