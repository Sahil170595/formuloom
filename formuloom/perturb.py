from __future__ import annotations

import datetime
import json
import logging
import os
import re
import uuid
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape as _xml_escape

import openpyxl  # type: ignore[import-untyped]
from openpyxl.formula.tokenizer import Token, Tokenizer, TokenizerError  # type: ignore[import-untyped]
from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries  # type: ignore[import-untyped]
from openpyxl.utils.datetime import to_excel as _to_excel_serial  # type: ignore[import-untyped]

from formuloom.bundle import (
    COMPLETE_FILE,
    INIT_FILE,
    INSTRUCTIONS_FILE,
    RAW_DIFF_FILE,
    SUBSET_FILE,
    MissingBundleFileError,
)
from formuloom.schema import DiffFile, cell_column_letters, cell_row

logger = logging.getLogger(__name__)

SCRATCH_BLOCK_ROWS = 3

SCRATCH_BLOCK_COLS = 3

SCRATCH_VALUE_BASE = 101

PAD_ROW_HEIGHT = 15.0

DEFAULT_PAD_ROWS = 5

DEFAULT_INSERT_ROWS = 2

ANCHOR_MARGIN_ROWS = 5

RENAME_SUFFIX = "_renamed"


class PerturbError(Exception):
    pass


class UnsupportedFormulaError(PerturbError):

    def __init__(self, offenders: list[tuple[str, str, str]]) -> None:
        self.offenders = offenders
        detail = "; ".join(f"{s}!{r}: {f}" for s, r, f in offenders)
        super().__init__(f"cannot safely row-shift {len(offenders)} formula(s): {detail}")


class RegionCollisionError(PerturbError):
    pass


class _NotARef(Exception):
    pass


class _UnshiftableRef(Exception):
    pass


@dataclass
class MutableBundle:

    source_dir: Path
    init_wb: Any
    complete_wb: Any
    complete_values_wb: Any
    raw_diff: dict[str, Any]
    subset: dict[str, Any] | None
    instructions: str
    init_formula_values: dict[str, dict[str, object]]
    complete_formula_values: dict[str, dict[str, object]]
    sheet_renames: dict[str, str] = field(default_factory=dict)
    row_shifts: dict[str, tuple[int, int]] = field(default_factory=dict)


def _capture_formula_values(formula_wb: Any, values_wb: Any) -> dict[str, dict[str, object]]:
    out: dict[str, dict[str, object]] = {}
    for sheet_name in formula_wb.sheetnames:
        fws = formula_wb[sheet_name]
        vws = values_wb[sheet_name]
        sheet_map: dict[str, object] = {}
        for row in fws.iter_rows():
            for cell in row:
                if _formula_text(cell.value) is None:
                    continue
                value = vws[cell.coordinate].value
                if value is not None:
                    sheet_map[cell.coordinate] = value
        out[sheet_name] = sheet_map
    return out


def load_bundle(bundle_dir: Path | str) -> MutableBundle:
    bundle_dir = Path(bundle_dir)
    init_path = bundle_dir / INIT_FILE
    complete_path = bundle_dir / COMPLETE_FILE
    raw_diff_path = bundle_dir / RAW_DIFF_FILE
    instructions_path = bundle_dir / INSTRUCTIONS_FILE
    subset_path = bundle_dir / SUBSET_FILE

    for required in (init_path, complete_path, raw_diff_path, instructions_path):
        if not required.exists():
            raise MissingBundleFileError(f"required bundle file missing: {required}")

    raw_diff = json.loads(raw_diff_path.read_text(encoding="utf-8"))
    DiffFile.model_validate(raw_diff)

    subset: dict[str, Any] | None = None
    if subset_path.exists():
        subset = json.loads(subset_path.read_text(encoding="utf-8"))
        DiffFile.model_validate(subset)

    init_wb = openpyxl.load_workbook(init_path, data_only=False)
    complete_wb = openpyxl.load_workbook(complete_path, data_only=False)
    init_values_wb = openpyxl.load_workbook(init_path, data_only=True)
    complete_values_wb = openpyxl.load_workbook(complete_path, data_only=True)

    return MutableBundle(
        source_dir=bundle_dir,
        init_wb=init_wb,
        complete_wb=complete_wb,
        complete_values_wb=complete_values_wb,
        raw_diff=raw_diff,
        subset=subset,
        instructions=instructions_path.read_text(encoding="utf-8"),
        init_formula_values=_capture_formula_values(init_wb, init_values_wb),
        complete_formula_values=_capture_formula_values(complete_wb, complete_values_wb),
    )


def _atomic_write_text(path: Path, text: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp-{os.getpid()}-{uuid.uuid4().hex}")
    try:
        tmp.write_text(text, encoding="utf-8")
        os.replace(tmp, path)
    except OSError:
        tmp.unlink(missing_ok=True)
        raise


def _atomic_write_workbook(wb: Any, path: Path, carryover: dict[str, dict[str, object]] | None = None) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_name(f".{path.name}.tmp-{os.getpid()}-{uuid.uuid4().hex}")
    try:
        wb.save(tmp)
        if carryover and any(carryover.values()):
            _patch_cached_values(tmp, carryover, epoch=wb.epoch)
        os.replace(tmp, path)
    except OSError:
        tmp.unlink(missing_ok=True)
        raise


def _remap_carryover(
    original: dict[str, dict[str, object]], sheet_renames: dict[str, str], row_shifts: dict[str, tuple[int, int]]
) -> dict[str, dict[str, object]]:
    out: dict[str, dict[str, object]] = {}
    for sheet_name, cell_map in original.items():
        new_sheet = sheet_renames.get(sheet_name, sheet_name)
        shift = row_shifts.get(sheet_name)
        remapped = out.setdefault(new_sheet, {})
        for ref, value in cell_map.items():
            new_ref = shift_cell_ref(ref, shift[0], shift[1]) if shift is not None else ref
            remapped[new_ref] = value
    return out


def write_bundle(bundle: MutableBundle, out_dir: Path | str) -> None:
    out_dir = Path(out_dir)
    DiffFile.model_validate(bundle.raw_diff)
    if bundle.subset is not None:
        DiffFile.model_validate(bundle.subset)

    init_carryover = _remap_carryover(bundle.init_formula_values, bundle.sheet_renames, bundle.row_shifts)
    complete_carryover = _remap_carryover(bundle.complete_formula_values, bundle.sheet_renames, bundle.row_shifts)
    _atomic_write_workbook(bundle.init_wb, out_dir / INIT_FILE, carryover=init_carryover)
    _atomic_write_workbook(bundle.complete_wb, out_dir / COMPLETE_FILE, carryover=complete_carryover)
    _atomic_write_text(out_dir / RAW_DIFF_FILE, json.dumps(bundle.raw_diff, indent=2, ensure_ascii=False) + "\n")
    _atomic_write_text(out_dir / INSTRUCTIONS_FILE, bundle.instructions)
    if bundle.subset is not None:
        _atomic_write_text(out_dir / SUBSET_FILE, json.dumps(bundle.subset, indent=2, ensure_ascii=False) + "\n")


_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"


def _sheet_name_to_part(zf: zipfile.ZipFile) -> dict[str, str]:
    wb_xml = ET.fromstring(zf.read("xl/workbook.xml"))
    rid_by_name = {
        el.get("name"): el.get(f"{{{_DOC_REL_NS}}}id")
        for el in wb_xml.iter(f"{{{_MAIN_NS}}}sheet")
        if el.get("name") is not None
    }
    rels_xml = ET.fromstring(zf.read("xl/_rels/workbook.xml.rels"))
    target_by_rid = {
        el.get("Id"): el.get("Target") for el in rels_xml.iter(f"{{{_PKG_REL_NS}}}Relationship") if el.get("Id")
    }
    out: dict[str, str] = {}
    for name, rid in rid_by_name.items():
        target = target_by_rid.get(rid)
        if name is None or target is None:
            continue
        out[name] = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
    return out


def _format_cached_value(value: object, epoch: Any) -> tuple[str, str | None] | None:
    if isinstance(value, bool):
        return ("1" if value else "0"), "b"
    if isinstance(value, int | float):
        return (repr(value) if isinstance(value, float) else str(value)), None
    if isinstance(value, datetime.datetime | datetime.date):
        return str(_to_excel_serial(value, epoch)), None
    if isinstance(value, str):
        return value, "str"
    return None


def _inject_cell_values(sheet_xml: str, values_for_sheet: dict[str, object], epoch: Any) -> str:
    if not values_for_sheet:
        return sheet_xml
    refs_alt = "|".join(re.escape(ref) for ref in values_for_sheet)
    cell_re = re.compile(r'<c r="(' + refs_alt + r')"((?:(?!/?>).)*)>((?:(?!</c>).)*)</c>', re.DOTALL)

    def _replace(m: re.Match[str]) -> str:
        ref, attrs, body = m.group(1), m.group(2), m.group(3)
        if "<f" not in body:
            return m.group(0)
        formatted = _format_cached_value(values_for_sheet[ref], epoch)
        if formatted is None:
            return m.group(0)
        v_text, t_attr = formatted
        new_body = re.sub(r"<v\s*/>|<v>\s*</v>", f"<v>{_xml_escape(v_text)}</v>", body, count=1)
        if new_body == body:
            return m.group(0)
        if t_attr is not None:
            attrs = f' t="{t_attr}"{attrs}'
        return f'<c r="{ref}"{attrs}>{new_body}</c>'

    return cell_re.sub(_replace, sheet_xml)


def _patch_cached_values(xlsx_path: Path, values_by_sheet: dict[str, dict[str, object]], epoch: Any) -> None:
    with zipfile.ZipFile(xlsx_path, "r") as zin:
        infos = zin.infolist()
        entries = {info.filename: zin.read(info.filename) for info in infos}
        sheet_parts = _sheet_name_to_part(zin)

    for sheet_name, values_for_sheet in values_by_sheet.items():
        if not values_for_sheet:
            continue
        part = sheet_parts.get(sheet_name)
        if part is None or part not in entries:
            logger.warning("cached-value carryover: sheet %r not found in %s", sheet_name, xlsx_path)
            continue
        xml_text = entries[part].decode("utf-8")
        entries[part] = _inject_cell_values(xml_text, values_for_sheet, epoch).encode("utf-8")

    tmp = xlsx_path.with_name(f".{xlsx_path.name}.patch-{uuid.uuid4().hex}")
    try:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
            for info in infos:
                zout.writestr(info, entries[info.filename])
        os.replace(tmp, xlsx_path)
    except OSError:
        tmp.unlink(missing_ok=True)
        raise


_QUOTED_SHEET_PREFIX_RE = re.compile(r"^'((?:[^']|'')*)'!")
_UNQUOTED_SHEET_PREFIX_RE = re.compile(r"^([A-Za-z_][A-Za-z0-9_.]*)!")
_SINGLE_CELL_RE = re.compile(r"^(\$?)([A-Za-z]{1,3})(\$?)([0-9]+)$")
_WHOLE_COL_RE = re.compile(r"^\$?[A-Za-z]{1,3}$")
_WHOLE_ROW_RE = re.compile(r"^\$?[0-9]+$")
_SHEET_NAME_OK_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_.]*$")


def _formula_text(value: object) -> str | None:
    if isinstance(value, str):
        return value if value.startswith("=") else None
    if value is not None and "Formula" in type(value).__name__:
        text = getattr(value, "text", None)
        if isinstance(text, str) and text:
            return text if text.startswith("=") else "=" + text
    return None


def _is_wrapper_formula(value: object) -> bool:
    return value is not None and not isinstance(value, str) and "Formula" in type(value).__name__


def _needs_quoting(sheet_name: str) -> bool:
    return _SHEET_NAME_OK_RE.match(sheet_name) is None


def _quote_sheet_name(sheet_name: str) -> str:
    if _needs_quoting(sheet_name):
        return "'" + sheet_name.replace("'", "''") + "'"
    return sheet_name


def _extract_sheet_prefix(token_value: str) -> tuple[str | None, int]:
    m = _QUOTED_SHEET_PREFIX_RE.match(token_value)
    if m is not None:
        return m.group(1).replace("''", "'"), m.end()
    m = _UNQUOTED_SHEET_PREFIX_RE.match(token_value)
    if m is not None:
        return m.group(1), m.end()
    return None, 0


def _rewrite_sheet_name_token(token_value: str, old: str, new: str) -> str:
    prefix, plen = _extract_sheet_prefix(token_value)
    if prefix != old:
        return token_value
    return f"{_quote_sheet_name(new)}!{token_value[plen:]}"


def _rewrite_formula_sheet_refs(formula: str, old: str, new: str) -> str:
    if old not in formula:
        return formula
    try:
        tk = Tokenizer(formula)
    except TokenizerError as exc:
        logger.warning("could not tokenize formula %r during sheet rename: %s", formula, exc)
        return formula
    changed = False
    for token in tk.items:
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            rewritten = _rewrite_sheet_name_token(token.value, old, new)
            if rewritten != token.value:
                token.value = rewritten
                changed = True
    return tk.render() if changed else formula


def _rewrite_ref_string_sheet(text: str, old: str, new: str) -> str:
    rewritten = _rewrite_formula_sheet_refs("=" + text, old, new)
    return rewritten[1:] if rewritten.startswith("=") else rewritten


def _shift_ref_body(ref_body: str, at_row: int, n: int) -> str:
    parts = ref_body.split(":")
    if len(parts) > 2:
        raise _UnshiftableRef(ref_body)
    shifted: list[str] = []
    for part in parts:
        m = _SINGLE_CELL_RE.match(part)
        if m is not None:
            dollar_col, col, dollar_row, row_str = m.groups()
            row = int(row_str)
            new_row = row + n if row >= at_row else row
            shifted.append(f"{dollar_col}{col}{dollar_row}{new_row}")
            continue
        if len(parts) == 2 and _WHOLE_COL_RE.match(part):
            shifted.append(part)
            continue
        if _WHOLE_ROW_RE.match(part):
            raise _UnshiftableRef(ref_body)
        raise _NotARef(ref_body)
    return ":".join(shifted)


def _classify_and_shift_range_token(
    token_value: str, target_sheet: str, current_sheet: str, at_row: int, n: int
) -> tuple[str, bool]:
    prefix, plen = _extract_sheet_prefix(token_value)
    if prefix is not None:
        if prefix != target_sheet:
            return token_value, False
        head, ref_body = token_value[:plen], token_value[plen:]
    else:
        if current_sheet != target_sheet:
            return token_value, False
        head, ref_body = "", token_value
    if "[" in ref_body:
        return token_value, True
    try:
        new_body = _shift_ref_body(ref_body, at_row, n)
    except _NotARef:
        return token_value, False
    except _UnshiftableRef:
        return token_value, True
    return head + new_body, False


def _rewrite_formula_for_insert(
    formula: str, target_sheet: str, current_sheet: str, at_row: int, n: int
) -> tuple[str, bool]:
    if target_sheet not in formula and current_sheet != target_sheet:
        return formula, False
    try:
        tk = Tokenizer(formula)
    except TokenizerError as exc:
        logger.warning("could not tokenize formula %r for row-insert: %s", formula, exc)
        return formula, True
    any_offender = False
    for token in tk.items:
        if token.type == Token.OPERAND and token.subtype == Token.RANGE:
            new_val, offender = _classify_and_shift_range_token(token.value, target_sheet, current_sheet, at_row, n)
            if offender:
                any_offender = True
            elif new_val != token.value:
                token.value = new_val
    if any_offender:
        return formula, True
    return tk.render(), False


def _collect_and_maybe_apply_row_shift(
    wb: Any, target_sheet: str, at_row: int, n: int, *, apply: bool
) -> list[tuple[str, str, str]]:
    offenders: list[tuple[str, str, str]] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for row in ws.iter_rows():
            for cell in row:
                text = _formula_text(cell.value)
                if text is None:
                    continue
                if _is_wrapper_formula(cell.value):

                    if sheet_name == target_sheet or target_sheet in text:
                        offenders.append((sheet_name, cell.coordinate, text))
                    continue
                new_text, hit_offender = _rewrite_formula_for_insert(text, target_sheet, sheet_name, at_row, n)
                if hit_offender:
                    offenders.append((sheet_name, cell.coordinate, text))
                elif apply and new_text != text:
                    cell.value = new_text
    return offenders


def shift_cell_ref(ref: str, at_row: int, delta: int) -> str:
    row = cell_row(ref)
    col = cell_column_letters(ref)
    new_row = row + delta if row >= at_row else row
    return f"{col}{new_row}"


def _assert_region_unused(ws: Any, start_row: int, start_col: int, nrows: int, ncols: int) -> None:
    for r in range(start_row, start_row + nrows):
        for c in range(start_col, start_col + ncols):
            if ws.cell(row=r, column=c).value is not None:
                raise RegionCollisionError(f"{ws.title}!{get_column_letter(c)}{r} is not empty — refusing to overwrite")


def add_scratch_block(bundle: MutableBundle, sheet: str, anchor: str) -> MutableBundle:
    ws = bundle.complete_wb[sheet]
    start_row = cell_row(anchor)
    start_col = column_index_from_string(cell_column_letters(anchor))
    _assert_region_unused(ws, start_row, start_col, SCRATCH_BLOCK_ROWS, SCRATCH_BLOCK_COLS)

    value_rows = SCRATCH_BLOCK_ROWS - 1
    for r in range(value_rows):
        for c in range(SCRATCH_BLOCK_COLS):
            ws.cell(row=start_row + r, column=start_col + c, value=SCRATCH_VALUE_BASE + r * SCRATCH_BLOCK_COLS + c)
    sum_row = start_row + value_rows
    first = f"{get_column_letter(start_col)}{start_row}"
    last = f"{get_column_letter(start_col + SCRATCH_BLOCK_COLS - 1)}{start_row + value_rows - 1}"
    ws.cell(row=sum_row, column=start_col, value=f"=SUM({first}:{last})")
    return bundle


def duplicate_region(bundle: MutableBundle, sheet: str, src_range: str, anchor: str) -> MutableBundle:
    min_col, min_row, max_col, max_row = range_boundaries(src_range)
    nrows, ncols = max_row - min_row + 1, max_col - min_col + 1
    dest_row = cell_row(anchor)
    dest_col = column_index_from_string(cell_column_letters(anchor))

    src_ws = bundle.complete_values_wb[sheet]
    values = [[src_ws.cell(row=min_row + r, column=min_col + c).value for c in range(ncols)] for r in range(nrows)]

    for wb in (bundle.init_wb, bundle.complete_wb):
        ws = wb[sheet]
        _assert_region_unused(ws, dest_row, dest_col, nrows, ncols)
        for r in range(nrows):
            for c in range(ncols):
                ws.cell(row=dest_row + r, column=dest_col + c, value=values[r][c])
    return bundle


def reorder_sheet_tabs(bundle: MutableBundle) -> MutableBundle:
    for wb in (bundle.init_wb, bundle.complete_wb):
        wb._sheets.reverse()
    return bundle


def pad_bottom_rows(bundle: MutableBundle, sheet: str, n: int) -> MutableBundle:
    if n <= 0:
        raise PerturbError(f"pad_bottom_rows: n must be positive, got {n}")
    bottom = max(bundle.init_wb[sheet].max_row, bundle.complete_wb[sheet].max_row)
    for wb in (bundle.init_wb, bundle.complete_wb):
        ws = wb[sheet]
        for r in range(bottom + 1, bottom + 1 + n):
            ws.row_dimensions[r].height = PAD_ROW_HEIGHT
    return bundle


def rename_sheet(bundle: MutableBundle, old: str, new: str) -> MutableBundle:
    if old == new:
        raise PerturbError(f"rename_sheet: old == new ({old!r})")

    for wb in (bundle.init_wb, bundle.complete_wb):
        if old not in wb.sheetnames:
            raise PerturbError(f"rename_sheet: sheet {old!r} not found in workbook")
        if new in wb.sheetnames:
            raise PerturbError(f"rename_sheet: target sheet name {new!r} already exists")

        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for row in ws.iter_rows():
                for cell in row:
                    text = _formula_text(cell.value)
                    if text is None:
                        continue
                    rewritten = _rewrite_formula_sheet_refs(text, old, new)
                    if rewritten == text:
                        continue
                    if _is_wrapper_formula(cell.value):
                        cell.value.text = rewritten
                    else:
                        cell.value = rewritten

        for name in list(wb.defined_names.keys()):
            dn = wb.defined_names[name]
            if dn.value:
                rewritten = _rewrite_ref_string_sheet(dn.value, old, new)
                if rewritten != dn.value:
                    dn.value = rewritten

        wb[old].title = new

    for diff in (bundle.raw_diff, bundle.subset):
        if diff is None:
            continue
        sheets = diff.get("sheets", {})
        if old in sheets:
            sheets[new] = sheets.pop(old)

    bundle.sheet_renames[old] = new
    return bundle


def insert_blank_rows(bundle: MutableBundle, sheet: str, at_row: int, n: int) -> MutableBundle:
    if n <= 0:
        raise PerturbError(f"insert_blank_rows: n must be positive, got {n}")
    if at_row <= 0:
        raise PerturbError(f"insert_blank_rows: at_row must be positive, got {at_row}")

    offenders: list[tuple[str, str, str]] = []
    for wb in (bundle.init_wb, bundle.complete_wb):
        if sheet not in wb.sheetnames:
            raise PerturbError(f"insert_blank_rows: sheet {sheet!r} not found in workbook")
        offenders.extend(_collect_and_maybe_apply_row_shift(wb, sheet, at_row, n, apply=False))
    if offenders:
        raise UnsupportedFormulaError(offenders)

    for wb in (bundle.init_wb, bundle.complete_wb):
        _collect_and_maybe_apply_row_shift(wb, sheet, at_row, n, apply=True)
        wb[sheet].insert_rows(at_row, amount=n)

    for diff in (bundle.raw_diff, bundle.subset):
        if diff is None:
            continue
        sheet_diff = diff.get("sheets", {}).get(sheet)
        if sheet_diff is None:
            continue
        for group_name in ("intermediate", "final"):
            group = sheet_diff.get("groups", {}).get(group_name)
            if not group:
                continue
            for entry in group.get("cells", []):
                entry["cell"] = shift_cell_ref(entry["cell"], at_row, n)

    bundle.row_shifts[sheet] = (at_row, n)
    return bundle


def _pick_target_sheet(bundle: MutableBundle) -> str:
    sheets: list[str] = list(bundle.raw_diff.get("sheets", {}).keys())
    if not sheets:
        raise PerturbError("perturb_suite: bundle raw_diff has no sheets")
    return sheets[0]


def _below_anchor(ws: Any, margin: int, col: str = "A") -> str:
    return f"{col}{ws.max_row + margin}"


def perturb_suite(bundle_dir: Path | str, out_root: Path | str) -> dict[str, Any]:
    bundle_dir = Path(bundle_dir)
    out_root = Path(out_root)
    manifest: dict[str, Any] = {"source": str(bundle_dir), "out_root": str(out_root), "perturbations": {}}

    probe = load_bundle(bundle_dir)
    target_sheet = _pick_target_sheet(probe)
    src_ws = probe.complete_wb[target_sheet]
    scratch_anchor = _below_anchor(src_ws, ANCHOR_MARGIN_ROWS)
    dup_anchor = _below_anchor(src_ws, ANCHOR_MARGIN_ROWS + SCRATCH_BLOCK_ROWS + 2)
    src_range = f"A1:{get_column_letter(min(2, src_ws.max_column))}{min(2, src_ws.max_row)}"
    insert_at_row = max(2, src_ws.max_row // 2)

    jobs: list[tuple[str, Callable[[MutableBundle], MutableBundle]]] = [
        ("add_scratch_block", lambda b: add_scratch_block(b, target_sheet, scratch_anchor)),
        ("duplicate_region", lambda b: duplicate_region(b, target_sheet, src_range, dup_anchor)),
        ("reorder_sheet_tabs", reorder_sheet_tabs),
        ("pad_bottom_rows", lambda b: pad_bottom_rows(b, target_sheet, DEFAULT_PAD_ROWS)),
        ("rename_sheet", lambda b: rename_sheet(b, target_sheet, target_sheet + RENAME_SUFFIX)),
        ("insert_blank_rows", lambda b: insert_blank_rows(b, target_sheet, insert_at_row, DEFAULT_INSERT_ROWS)),
    ]

    for name, fn in jobs:
        entry: dict[str, Any] = {"sheet": target_sheet}
        try:
            fresh = load_bundle(bundle_dir)
            fn(fresh)
            out_dir = out_root / name
            write_bundle(fresh, out_dir)
            entry["status"] = "ok"
            entry["out_dir"] = str(out_dir)
        except UnsupportedFormulaError as exc:
            logger.warning("perturb_suite: %s unsupported on %s: %s", name, bundle_dir, exc)
            entry["status"] = "unsupported"
            entry["reason"] = str(exc)
            entry["offenders"] = [{"sheet": s, "ref": r, "formula": f} for s, r, f in exc.offenders]
        except PerturbError as exc:
            logger.warning("perturb_suite: %s failed on %s: %s", name, bundle_dir, exc)
            entry["status"] = "error"
            entry["reason"] = str(exc)
        manifest["perturbations"][name] = entry

    return manifest
