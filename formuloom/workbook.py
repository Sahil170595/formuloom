from __future__ import annotations

import colorsys
import hashlib
import logging
import pickle
import re
import warnings
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]
from openpyxl.styles.colors import COLOR_INDEX  # type: ignore[import-untyped]
from openpyxl.utils import get_column_letter  # type: ignore[import-untyped]

logger = logging.getLogger(__name__)

CACHE_VERSION = "wb-v1"

CACHE_DIR_DEFAULT = Path(".cache")

CACHE_SUBDIR = "workbook"

SHA256_CHUNK_BYTES = 1 << 20

INPUT_BLUE_CHANNEL_MARGIN = 40

CLASSIC_INPUT_RGBS = frozenset({"0000FF", "1F4E79", "1F497D", "0070C0", "00205C"})

_THEME_XML_ELEMENTS = (
    "dk1",
    "lt1",
    "dk2",
    "lt2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
)
_THEME_INDEX_TO_ELEMENT = (
    "lt1",
    "dk1",
    "lt2",
    "dk2",
    "accent1",
    "accent2",
    "accent3",
    "accent4",
    "accent5",
    "accent6",
    "hlink",
    "folHlink",
)


@dataclass(frozen=True)
class CellRecord:

    sheet: str
    ref: str
    row: int
    col: int
    value: object | None
    formula: str | None
    font_rgb: str | None
    is_input_colored: bool
    bold: bool
    border_top: bool
    border_bottom: bool
    number_format: str
    merged_anchor: str | None


@dataclass
class SheetData:

    name: str
    cells: dict[str, CellRecord]
    max_row: int
    max_col: int
    hidden: bool
    frozen_panes: str | None
    hidden_rows: set[int]
    hidden_cols: set[str]


@dataclass
class WorkbookData:

    path: Path
    sha256: str
    sheets: dict[str, SheetData]
    defined_names: dict[str, str]
    chart_source_refs: dict[str, list[str]]
    print_areas: dict[str, str | None]
    table_ranges: dict[str, list[str]]


def _hex6(argb: str) -> str:
    s = argb.strip().upper()
    if len(s) == 8:
        s = s[2:]
    return s


def parse_theme_colors(theme_bytes: bytes | None) -> list[str]:
    if not theme_bytes:
        return []
    text = theme_bytes.decode("utf-8", errors="replace")
    scheme = re.search(r"<a:clrScheme\b.*?</a:clrScheme>", text, re.DOTALL)
    if scheme is None:
        return []
    block = scheme.group(0)
    by_element: dict[str, str] = {}
    for elem in _THEME_XML_ELEMENTS:
        seg_match = re.search(rf"<a:{elem}\b.*?</a:{elem}>", block, re.DOTALL)
        if seg_match is None:
            continue
        seg = seg_match.group(0)
        srgb = re.search(r'<a:srgbClr\s+val="([0-9A-Fa-f]{6})"', seg)
        if srgb is not None:
            by_element[elem] = srgb.group(1).upper()
            continue
        sysc = re.search(r'<a:sysClr\b[^>]*lastClr="([0-9A-Fa-f]{6})"', seg)
        if sysc is not None:
            by_element[elem] = sysc.group(1).upper()
    return [by_element.get(el, "") for el in _THEME_INDEX_TO_ELEMENT]


def apply_tint(rgb6: str, tint: float) -> str:
    if not tint:
        return rgb6.upper()
    r = int(rgb6[0:2], 16) / 255.0
    g = int(rgb6[2:4], 16) / 255.0
    b = int(rgb6[4:6], 16) / 255.0
    hue, lum, sat = colorsys.rgb_to_hls(r, g, b)
    lum = lum * (1.0 + tint) if tint < 0 else lum * (1.0 - tint) + tint
    lum = min(1.0, max(0.0, lum))
    nr, ng, nb = colorsys.hls_to_rgb(hue, lum, sat)
    return f"{round(nr * 255):02X}{round(ng * 255):02X}{round(nb * 255):02X}"


def resolve_font_rgb(color: Any, theme_colors: list[str]) -> str | None:
    if color is None:
        return None
    ctype = getattr(color, "type", None)
    tint = float(getattr(color, "tint", 0.0) or 0.0)
    if ctype == "rgb":
        raw = getattr(color, "rgb", None)
        return _hex6(raw) if isinstance(raw, str) else None
    if ctype == "theme":
        idx = getattr(color, "theme", None)
        if not isinstance(idx, int) or not 0 <= idx < len(theme_colors):
            return None
        base = theme_colors[idx]
        return apply_tint(base, tint) if base else None
    if ctype == "indexed":
        idx = getattr(color, "indexed", None)
        if not isinstance(idx, int) or not 0 <= idx < len(COLOR_INDEX):
            return None
        return apply_tint(_hex6(str(COLOR_INDEX[idx])), tint)
    return None


def is_blue_input(rgb6: str | None) -> bool:
    if rgb6 is None or len(rgb6) != 6:
        return False
    up = rgb6.upper()
    if up in CLASSIC_INPUT_RGBS:
        return True
    try:
        red, green, blue = int(up[0:2], 16), int(up[2:4], 16), int(up[4:6], 16)
    except ValueError:
        return False
    return blue > red + INPUT_BLUE_CHANNEL_MARGIN and blue > green + INPUT_BLUE_CHANNEL_MARGIN


def _extract_chart_refs(ws: Any) -> list[str]:
    refs: list[str] = []
    charts = getattr(ws, "_charts", None)
    if not charts:
        return refs
    for chart in charts:
        for series in getattr(chart, "series", None) or []:
            for axis_attr in ("val", "cat"):
                axis = getattr(series, axis_attr, None)
                if axis is None:
                    continue
                for ref_attr in ("numRef", "strRef"):
                    ref = getattr(axis, ref_attr, None)
                    formula = getattr(ref, "f", None) if ref is not None else None
                    if isinstance(formula, str) and formula:
                        refs.append(formula)
    return refs


def _extract_print_area(ws: Any) -> str | None:
    area = ws.print_area
    if not area:
        return None
    if isinstance(area, (list, tuple)):
        return ",".join(str(a) for a in area) or None
    return str(area)


def _extract_table_ranges(ws: Any) -> list[str]:
    out: list[str] = []
    for table in dict(ws.tables).values():
        ref = getattr(table, "ref", None)
        if ref:
            out.append(str(ref))
    return out


def _extract_defined_names(wb: Any) -> dict[str, str]:
    out: dict[str, str] = {}
    for name in wb.defined_names:
        dn = wb.defined_names[name]
        dest = getattr(dn, "value", None)
        out[str(name)] = str(dest) if dest is not None else ""
    return out


def _safe(label: str, path: Path, fn: Any) -> Any:
    try:
        return fn()
    except Exception as exc:  # noqa: BLE001 - presentation extraction is best-effort and logged
        logger.warning("presentation signal %r failed for %s: %s", label, path, exc)
        return None


def _formula_text(fval: object) -> str | None:
    if isinstance(fval, str):
        return fval if fval.startswith("=") else None
    if fval is not None and "Formula" in type(fval).__name__:
        text = getattr(fval, "text", None)
        if isinstance(text, str) and text:
            return text if text.startswith("=") else "=" + text
    return None


def _side_present(side: Any) -> bool:
    return bool(side is not None and getattr(side, "style", None))


def _build_merged_anchor_map(fws: Any) -> dict[str, str]:
    out: dict[str, str] = {}
    for cell_range in fws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = cell_range.bounds
        anchor = f"{get_column_letter(min_col)}{min_row}"
        for rr in range(min_row, max_row + 1):
            for cc in range(min_col, max_col + 1):
                out[f"{get_column_letter(cc)}{rr}"] = anchor
    return out


def _extract_sheet(sheet_name: str, fws: Any, vws: Any, theme_colors: list[str]) -> SheetData:
    merged = _build_merged_anchor_map(fws)
    cells: dict[str, CellRecord] = {}
    for row in fws.iter_rows():
        for fcell in row:
            fval = fcell.value
            if fval is None:
                continue
            ref = fcell.coordinate
            value = vws[ref].value
            font = fcell.font
            color = getattr(font, "color", None) if font is not None else None
            rgb = resolve_font_rgb(color, theme_colors)
            border = fcell.border
            cells[ref] = CellRecord(
                sheet=sheet_name,
                ref=ref,
                row=int(fcell.row),
                col=int(fcell.column),
                value=value,
                formula=_formula_text(fval),
                font_rgb=rgb,
                is_input_colored=is_blue_input(rgb),
                bold=bool(getattr(font, "bold", False)) if font is not None else False,
                border_top=_side_present(getattr(border, "top", None)) if border is not None else False,
                border_bottom=_side_present(getattr(border, "bottom", None)) if border is not None else False,
                number_format=str(fcell.number_format),
                merged_anchor=merged.get(ref),
            )
    hidden_rows = {int(r) for r, dim in fws.row_dimensions.items() if getattr(dim, "hidden", False)}
    hidden_cols = {str(c) for c, dim in fws.column_dimensions.items() if getattr(dim, "hidden", False)}
    frozen = fws.freeze_panes
    return SheetData(
        name=sheet_name,
        cells=cells,
        max_row=int(fws.max_row or 0),
        max_col=int(fws.max_column or 0),
        hidden=(fws.sheet_state != "visible"),
        frozen_panes=str(frozen) if frozen else None,
        hidden_rows=hidden_rows,
        hidden_cols=hidden_cols,
    )


_PARSE_COUNT = 0


def get_parse_count() -> int:
    return _PARSE_COUNT


def reset_parse_count() -> None:
    global _PARSE_COUNT
    _PARSE_COUNT = 0


def sha256_of_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(SHA256_CHUNK_BYTES), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load_workbook_recording_warnings(path: Path, *, data_only: bool) -> Any:
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        wb = openpyxl.load_workbook(path, data_only=data_only, read_only=False)
    for w in caught:
        logger.warning("openpyxl warning loading %s (data_only=%s): %s", path, data_only, w.message)
    return wb


def _extract_workbook(path: Path, sha: str) -> WorkbookData:
    global _PARSE_COUNT
    _PARSE_COUNT += 1
    fwb = _load_workbook_recording_warnings(path, data_only=False)
    vwb = _load_workbook_recording_warnings(path, data_only=True)
    theme_colors = _safe("theme", path, lambda: parse_theme_colors(fwb.loaded_theme)) or []

    sheets: dict[str, SheetData] = {}
    chart_source_refs: dict[str, list[str]] = {}
    print_areas: dict[str, str | None] = {}
    table_ranges: dict[str, list[str]] = {}
    for sheet_name in fwb.sheetnames:
        fws = fwb[sheet_name]
        vws = vwb[sheet_name]
        sheets[sheet_name] = _extract_sheet(sheet_name, fws, vws, theme_colors)
        chart_source_refs[sheet_name] = _safe("charts", path, lambda ws=fws: _extract_chart_refs(ws)) or []
        print_areas[sheet_name] = _safe("print_area", path, lambda ws=fws: _extract_print_area(ws))
        table_ranges[sheet_name] = _safe("tables", path, lambda ws=fws: _extract_table_ranges(ws)) or []

    defined_names = _safe("defined_names", path, lambda: _extract_defined_names(fwb)) or {}
    return WorkbookData(
        path=path,
        sha256=sha,
        sheets=sheets,
        defined_names=defined_names,
        chart_source_refs=chart_source_refs,
        print_areas=print_areas,
        table_ranges=table_ranges,
    )


def _cache_path(cache_dir: Path, sha: str) -> Path:
    return cache_dir / CACHE_SUBDIR / f"{CACHE_VERSION}-{sha}.pkl"


def load_workbook_data(path: Path, cache_dir: Path | None = None) -> WorkbookData:
    path = Path(path)
    root = CACHE_DIR_DEFAULT if cache_dir is None else Path(cache_dir)
    sha = sha256_of_file(path)
    cpath = _cache_path(root, sha)
    if cpath.is_file():
        cached = _read_cache(cpath)
        if cached is not None:
            return cached
    data = _extract_workbook(path, sha)
    _write_cache(cpath, data)
    return data


def _read_cache(cpath: Path) -> WorkbookData | None:
    try:
        with cpath.open("rb") as fh:
            obj = pickle.load(fh)
    except Exception as exc:  # noqa: BLE001 - corrupt/partial cache -> re-parse, logged
        logger.warning("ignoring unreadable cache %s: %s", cpath, exc)
        return None
    if isinstance(obj, WorkbookData):
        return obj
    logger.warning("ignoring cache %s: unexpected object %s", cpath, type(obj))
    return None


def _write_cache(cpath: Path, data: WorkbookData) -> None:
    try:
        cpath.parent.mkdir(parents=True, exist_ok=True)
        with cpath.open("wb") as fh:
            pickle.dump(data, fh, protocol=pickle.HIGHEST_PROTOCOL)
    except OSError as exc:
        logger.warning("could not write cache %s: %s", cpath, exc)


@dataclass
class WorkbookPair:

    init: WorkbookData
    complete: WorkbookData

    @classmethod
    def load(cls, init_path: Path, complete_path: Path, cache_dir: Path | None = None) -> WorkbookPair:
        return cls(
            init=load_workbook_data(Path(init_path), cache_dir),
            complete=load_workbook_data(Path(complete_path), cache_dir),
        )


def value_diff(pair: WorkbookPair) -> dict[str, set[str]]:
    out: dict[str, set[str]] = {}
    init_sheets = pair.init.sheets
    for sheet_name, csheet in pair.complete.sheets.items():
        isheet = init_sheets.get(sheet_name)
        icells = isheet.cells if isheet is not None else {}
        changed: set[str] = set()
        for ref in set(csheet.cells) | set(icells):
            crec = csheet.cells.get(ref)
            irec = icells.get(ref)
            cval = crec.value if crec is not None else None
            ival = irec.value if irec is not None else None
            if cval != ival and not (cval is None and ival is None):
                changed.add(ref)
        out[sheet_name] = changed
    return out
