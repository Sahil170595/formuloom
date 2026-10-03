from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import openpyxl  # type: ignore[import-untyped]
import pytest
from openai import AsyncOpenAI
from openai.types.responses import Response
from typer.testing import CliRunner

from formuloom.bundle import TaskBundle
from formuloom.cli import app, predict_v0
from formuloom.schema import DiffFile

runner = CliRunner()

DATA = Path(__file__).parent.parent / "data"
requires_data = pytest.mark.skipif(not DATA.is_dir(), reason="data/ bundles not present")

_RAW_DIFF: dict[str, Any] = {
    "spec_version": 2,
    "task_tolerance": 0.01,
    "sheets": {
        "S1": {
            "sheet_weight": 1.0,
            "groups": {
                "intermediate": {
                    "weight": 1,
                    "cells": [
                        {"cell": "B1", "cell_type": "currency"},
                        {"cell": "B2", "cell_type": "currency"},
                        {"cell": "B3", "cell_type": "currency"},
                    ],
                }
            },
        }
    },
}

_GOLDEN: dict[str, Any] = {
    "spec_version": 2,
    "task_tolerance": 0.01,
    "sheets": {
        "S1": {
            "sheet_weight": 1.0,
            "groups": {
                "intermediate": {
                    "weight": 0.6,
                    "cells": [{"cell": "B2", "cell_type": "currency"}, {"cell": "B3", "cell_type": "currency"}],
                },
                "final": {"weight": 0.4, "cells": [{"cell": "B1", "cell_type": "currency"}]},
            },
        }
    },
}


def _write_fixture_workbook(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S1"
    ws["A1"] = "Total Revenue"
    ws["B1"] = "=SUM(B2:B3)"
    ws["A2"] = "Product A Sales"
    ws["B2"] = 100
    ws["A3"] = "Product B Sales"
    ws["B3"] = 50
    wb.save(path)


def _write_task_bundle(root: Path, *, with_subset: bool = True) -> Path:
    root.mkdir(parents=True, exist_ok=True)
    _write_fixture_workbook(root / "init.xlsx")
    _write_fixture_workbook(root / "complete.xlsx")
    (root / "instructions.md").write_text("Build a simple revenue rollup model.", encoding="utf-8")
    (root / "raw_diff.json").write_text(json.dumps(_RAW_DIFF), encoding="utf-8")
    if with_subset:
        (root / "subset.json").write_text(json.dumps(_GOLDEN), encoding="utf-8")
    return root


def _usage_payload(input_tokens: int = 100, output_tokens: int = 20) -> dict[str, Any]:
    return {
        "input_tokens": input_tokens,
        "input_tokens_details": {"cached_tokens": 0},
        "output_tokens": output_tokens,
        "output_tokens_details": {"reasoning_tokens": 0},
        "total_tokens": input_tokens + output_tokens,
    }


def _message_response(text: str, response_id: str = "resp") -> Response:
    payload = {
        "id": response_id,
        "created_at": 1.0,
        "model": "gpt-5.4-mini",
        "object": "response",
        "output": [
            {
                "id": "msg_1",
                "type": "message",
                "role": "assistant",
                "status": "completed",
                "content": [{"type": "output_text", "text": text, "annotations": []}],
            }
        ],
        "parallel_tool_calls": True,
        "tool_choice": "auto",
        "tools": [],
        "status": "completed",
        "usage": _usage_payload(),
    }
    return Response.model_validate(payload)


def _incomplete_response(reason: str = "max_output_tokens") -> Response:
    payload = {
        "id": "resp_incomplete",
        "created_at": 1.0,
        "model": "gpt-5.4-mini",
        "object": "response",
        "output": [],
        "parallel_tool_calls": True,
        "tool_choice": "auto",
        "tools": [],
        "status": "incomplete",
        "incomplete_details": {"reason": reason},
        "usage": _usage_payload(50, 5),
    }
    return Response.model_validate(payload)


_TASK_PROFILE_RESPONSE = json.dumps(
    {
        "workbook_purpose": "revenue rollup",
        "model_type": "3-statement",
        "expected_deliverables": ["total revenue"],
        "likely_capstone_outputs": ["Total Revenue"],
    }
)


def _install_fake_client(monkeypatch: pytest.MonkeyPatch, steps: list[Response]) -> list[dict[str, Any]]:
    calls: list[dict[str, Any]] = []
    remaining = list(steps)

    async def fake_create(**kwargs: Any) -> Response:
        calls.append(kwargs)
        return remaining.pop(0)

    client = AsyncOpenAI(api_key="test-key", max_retries=0)
    client.responses.create = fake_create  # type: ignore[assignment]
    monkeypatch.setattr("formuloom.cli.build_async_client", lambda settings: client)
    return calls


def test_predict_v0_surfaces_the_aggregating_labeled_row_as_final(tmp_path: Path) -> None:
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    bundle = TaskBundle.load(bundle_dir, "predict")
    result = predict_v0(bundle, cache_dir=tmp_path / ".cache")
    assert result.final_refs("S1") == {"B1"}
    assert result.intermediate_refs("S1") == {"B2", "B3"}


def test_predict_v0_output_round_trips_as_diff_file(tmp_path: Path) -> None:
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    bundle = TaskBundle.load(bundle_dir, "predict")
    result = predict_v0(bundle, cache_dir=tmp_path / ".cache")
    reloaded = DiffFile.model_validate(result.model_dump())
    assert reloaded == result


def test_predict_v0_never_touches_golden(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    bundle = TaskBundle.load(bundle_dir, "predict")
    assert bundle.golden is None
    result = predict_v0(bundle, cache_dir=tmp_path / ".cache")
    assert result.final_refs("S1") == {"B1"}


def test_cli_predict_v0_writes_schema_valid_generated_json(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    out_path = tmp_path / "out" / "generated.json"
    result = runner.invoke(app, ["predict", str(bundle_dir), "--variant", "V0", "-o", str(out_path)])
    assert result.exit_code == 0, result.output
    assert out_path.is_file()
    diff = DiffFile.model_validate(json.loads(out_path.read_text(encoding="utf-8")))
    assert diff.final_refs("S1") == {"B1"}


def test_cli_predict_v0_default_output_lands_in_run_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    result = runner.invoke(app, ["predict", str(bundle_dir), "--variant", "V0"])
    assert result.exit_code == 0, result.output
    generated = list((tmp_path / "runs").glob("*-V0/generated.json"))
    assert len(generated) == 1


def test_cli_predict_unknown_variant_fails_cleanly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    result = runner.invoke(app, ["predict", str(bundle_dir), "--variant", "V99"])
    assert result.exit_code == 1
    assert "unknown variant" in result.output


def test_cli_predict_v1_llm_variant_hermetic_end_to_end(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    row_response = json.dumps({"notes": "row 1 aggregates and is labeled Total", "final_rows": [1]})
    calls = _install_fake_client(
        monkeypatch, [_message_response(_TASK_PROFILE_RESPONSE), _message_response(row_response)]
    )

    out_path = tmp_path / "out" / "generated.json"
    result = runner.invoke(app, ["predict", str(bundle_dir), "--variant", "V1", "-o", str(out_path)])
    assert result.exit_code == 0, result.output
    assert len(calls) == 2

    diff = DiffFile.model_validate(json.loads(out_path.read_text(encoding="utf-8")))
    assert diff.final_refs("S1") == {"B1"}


def test_cli_predict_v1_best_effort_degrades_failed_sheet_and_exits_nonzero(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")

    _install_fake_client(monkeypatch, [_message_response(_TASK_PROFILE_RESPONSE), _incomplete_response()])

    out_path = tmp_path / "out" / "generated.json"
    result = runner.invoke(app, ["predict", str(bundle_dir), "--variant", "V1", "-o", str(out_path), "--best-effort"])
    assert result.exit_code == 1
    assert "degraded" in result.output

    diff = DiffFile.model_validate(json.loads(out_path.read_text(encoding="utf-8")))
    assert diff.final_refs("S1") == set()
    assert diff.intermediate_refs("S1") == {"B1", "B2", "B3"}

    failures_path = list((tmp_path / "out").parent.glob("**/failures.json"))

    run_failures = list((tmp_path / "runs").glob("*-V1/failures.json"))
    assert len(run_failures) == 1
    manifest = json.loads(run_failures[0].read_text(encoding="utf-8"))
    assert manifest == {"S1": manifest["S1"]}
    del failures_path


def test_cli_predict_v1_without_best_effort_hard_fails_on_sheet_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    _install_fake_client(monkeypatch, [_message_response(_TASK_PROFILE_RESPONSE), _incomplete_response()])

    out_path = tmp_path / "out" / "generated.json"
    result = runner.invoke(app, ["predict", str(bundle_dir), "--variant", "V1", "-o", str(out_path)])
    assert result.exit_code == 1
    assert not out_path.exists()


def test_cli_predict_nonexistent_task_dir_fails() -> None:
    result = runner.invoke(app, ["predict", "no/such/dir", "--variant", "V0"])
    assert result.exit_code != 0


def test_cli_score_both_modes_prints_compat_before_strict(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    pred_path = bundle_dir / "generated.json"
    pred_path.write_text(json.dumps(_GOLDEN), encoding="utf-8")

    result = runner.invoke(app, ["score", str(bundle_dir)])
    assert result.exit_code == 0, result.output
    compat_idx = result.output.index("[annotated]")
    strict_idx = result.output.index("[strict]")
    assert compat_idx < strict_idx
    assert "PRIMARY" in result.output.splitlines()[1]

    line = next(line for line in result.output.splitlines() if "[annotated]" in line)
    assert line.index("recall=") < line.index("precision=") < line.index("f1=")


def test_cli_score_oracle_prediction_is_perfect(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    (bundle_dir / "generated.json").write_text(json.dumps(_GOLDEN), encoding="utf-8")

    result = runner.invoke(app, ["score", str(bundle_dir), "--mode", "compat"])
    assert result.exit_code == 0, result.output
    assert "recall=1.0000" in result.output
    assert "precision=1.0000" in result.output
    assert "f1=1.0000" in result.output


def test_cli_score_missing_pred_file_fails_cleanly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    result = runner.invoke(app, ["score", str(bundle_dir)])
    assert result.exit_code == 1
    assert "prediction file not found" in result.output


def test_cli_score_explicit_pred_path(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    pred_path = tmp_path / "elsewhere" / "pred.json"
    pred_path.parent.mkdir(parents=True)
    pred_path.write_text(json.dumps(_GOLDEN), encoding="utf-8")
    result = runner.invoke(app, ["score", str(bundle_dir), "--pred", str(pred_path), "--mode", "strict"])
    assert result.exit_code == 0, result.output
    assert "[strict]" in result.output


def test_cli_score_bad_mode_fails_cleanly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    (bundle_dir / "generated.json").write_text(json.dumps(_GOLDEN), encoding="utf-8")
    result = runner.invoke(app, ["score", str(bundle_dir), "--mode", "bogus"])
    assert result.exit_code == 1
    assert "unknown --mode" in result.output


def test_cli_score_no_subset_fails_cleanly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture", with_subset=False)
    result = runner.invoke(app, ["score", str(bundle_dir)])
    assert result.exit_code == 1


def test_cli_eval_v0_offline_writes_eval_results_and_error_reports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    data_root = tmp_path / "data"
    _write_task_bundle(data_root / "Fixture")
    out_path = tmp_path / "eval_results.json"

    result = runner.invoke(app, ["eval", str(data_root), "--variant", "V0", "--offline-only", "--out", str(out_path)])
    assert result.exit_code == 0, result.output
    payload = json.loads(out_path.read_text(encoding="utf-8"))
    assert set(payload.keys()) == {"Fixture", "meta"}
    assert set(payload["Fixture"].keys()) == {"precision", "recall", "f1"}
    assert payload["Fixture"]["precision"] == 1.0
    assert payload["meta"]["mode"] == "annotated"
    assert "strict" in payload["meta"]
    assert payload["meta"]["variant"] == "V0"
    assert payload["meta"]["tokens"] == 0
    assert payload["meta"]["cost_usd"] == 0.0

    run_dir = out_path.parent

    generated = list((tmp_path / "runs").glob("*-V0/predictions/Fixture.generated.json"))
    errors = list((tmp_path / "runs").glob("*-V0/errors/Fixture.json"))
    assert len(generated) == 1
    assert len(errors) == 1
    del run_dir


def test_cli_eval_v0_default_out_lands_in_run_dir(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    data_root = tmp_path / "data"
    _write_task_bundle(data_root / "Fixture")
    result = runner.invoke(app, ["eval", str(data_root), "--variant", "V0"])
    assert result.exit_code == 0, result.output
    written = list((tmp_path / "runs").glob("*-V0/eval_results.json"))
    assert len(written) == 1


def test_cli_eval_repeat_records_repeat_meta(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    data_root = tmp_path / "data"
    _write_task_bundle(data_root / "Fixture")
    out_path = tmp_path / "eval_results.json"
    result = runner.invoke(app, ["eval", str(data_root), "--variant", "V0", "--repeat", "3", "--out", str(out_path)])
    assert result.exit_code == 0, result.output
    payload = json.loads(out_path.read_text(encoding="utf-8"))
    repeat_meta = payload["meta"]["repeat"]
    assert repeat_meta["n"] == 3
    assert len(repeat_meta["compat_micro_f1_per_run"]) == 3

    assert repeat_meta["stddev"] == pytest.approx(0.0)


def test_cli_eval_repeat_zero_rejected(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    data_root = tmp_path / "data"
    _write_task_bundle(data_root / "Fixture")
    result = runner.invoke(app, ["eval", str(data_root), "--variant", "V0", "--repeat", "0"])
    assert result.exit_code == 1
    assert "--repeat must be >= 1" in result.output


def test_cli_eval_v1_llm_variant_hermetic_end_to_end(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    data_root = tmp_path / "data"
    _write_task_bundle(data_root / "Fixture")
    row_response = json.dumps({"notes": "row 1 is the total", "final_rows": [1]})
    calls = _install_fake_client(
        monkeypatch, [_message_response(_TASK_PROFILE_RESPONSE), _message_response(row_response)]
    )

    out_path = tmp_path / "eval_results.json"
    result = runner.invoke(app, ["eval", str(data_root), "--variant", "V1", "--out", str(out_path)])
    assert result.exit_code == 0, result.output
    assert len(calls) == 2

    payload = json.loads(out_path.read_text(encoding="utf-8"))
    assert payload["Fixture"]["precision"] == 1.0
    assert payload["meta"]["variant"] == "V1"
    assert payload["meta"]["tokens"] > 0


def test_cli_predict_v4_selective_voting_triggers_on_v0_disagreement(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    probe = json.dumps({"notes": "rows 1 and 2 look final", "final_rows": [1, 2]})
    vote = json.dumps({"notes": "only row 1 is final", "final_rows": [1]})
    calls = _install_fake_client(
        monkeypatch,
        [_message_response(_TASK_PROFILE_RESPONSE), _message_response(probe), *([_message_response(vote)] * 3)],
    )

    out_path = tmp_path / "out" / "generated.json"
    result = runner.invoke(app, ["predict", str(bundle_dir), "--variant", "V4", "-o", str(out_path)])
    assert result.exit_code == 0, result.output

    assert len(calls) == 5

    diff = DiffFile.model_validate(json.loads(out_path.read_text(encoding="utf-8")))
    assert diff.final_refs("S1") == {"B1"}


def test_cli_predict_v12_uses_diverse_prompt_ensemble(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    bundle_dir = _write_task_bundle(tmp_path / "Fixture")
    row_response = json.dumps({"notes": "row 1 is the total", "final_rows": [1]})
    calls = _install_fake_client(
        monkeypatch,
        [
            _message_response(_TASK_PROFILE_RESPONSE),
            _message_response(row_response),
            _message_response(row_response),
            _message_response(row_response),
        ],
    )

    out_path = tmp_path / "out" / "generated.json"
    result = runner.invoke(app, ["predict", str(bundle_dir), "--variant", "V12", "-o", str(out_path)])
    assert result.exit_code == 0, result.output
    assert len(calls) == 4

    diff = DiffFile.model_validate(json.loads(out_path.read_text(encoding="utf-8")))
    assert diff.final_refs("S1") == {"B1"}


def test_cli_eval_llm_variant_with_offline_only_refuses_cleanly(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    data_root = tmp_path / "data"
    _write_task_bundle(data_root / "Fixture")
    result = runner.invoke(app, ["eval", str(data_root), "--variant", "V1", "--offline-only"])
    assert result.exit_code == 1
    assert "offline-only" in result.output.replace("_", "-")


def test_cli_eval_empty_data_root_fails_cleanly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    empty_root = tmp_path / "empty"
    empty_root.mkdir()
    result = runner.invoke(app, ["eval", str(empty_root), "--variant", "V0"])
    assert result.exit_code == 1
    assert "no task bundles found" in result.output


def test_cli_compare_renders_matrix_from_two_runs(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    data_root = tmp_path / "data"
    _write_task_bundle(data_root / "Fixture")

    run_a = tmp_path / "runs" / "run-a"
    run_b = tmp_path / "runs" / "run-b"
    eval_a = runner.invoke(app, ["eval", str(data_root), "--variant", "V0", "--out", str(run_a / "eval_results.json")])
    eval_b = runner.invoke(app, ["eval", str(data_root), "--variant", "V0", "--out", str(run_b / "eval_results.json")])
    assert eval_a.exit_code == 0, eval_a.output
    assert eval_b.exit_code == 0, eval_b.output

    result = runner.invoke(app, ["compare", str(run_a), str(run_b)])
    assert result.exit_code == 0, result.output
    assert "run-a" in result.output
    assert "run-b" in result.output
    assert "Fixture" in result.output
    assert "recall" in result.output and "precision" in result.output and "f1" in result.output
    assert "(strict)" in result.output
    assert "variant=V0" in result.output


def test_cli_compare_missing_eval_results_fails_cleanly(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.chdir(tmp_path)
    empty_run = tmp_path / "runs" / "nothing-here"
    empty_run.mkdir(parents=True)
    result = runner.invoke(app, ["compare", str(empty_run)])
    assert result.exit_code == 1
    assert "no eval_results.json" in result.output


def test_cli_version_flag_still_works_with_subcommands_registered() -> None:
    result = runner.invoke(app, ["--version"])
    assert result.exit_code == 0
