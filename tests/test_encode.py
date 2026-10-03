from __future__ import annotations

import re
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]
import pytest
from openpyxl.styles import Font  # type: ignore[import-untyped]

from formuloom.bundle import TaskBundle
from formuloom.encode import (
    SheetContext,
    TaskProfile,
    choose_grouping,
    estimate_tokens,
    split_by_sections,
)
from formuloom.features import TaskFeatures, build_task_features
from formuloom.schema import DiffFile, VariantConfig, cell_row
from formuloom.workbook import WorkbookData, load_workbook_data

DATA_ROOT = Path(__file__).resolve().parents[1] / "data"
TASKS = ["synthetic-statement", "synthetic-budget", "synthetic-rollup", "synthetic-projection", "synthetic-summary"]

ROW_LINE_RE = re.compile(r"^row (\d+) \|", re.MULTILINE)
CELL_LINE_RE = re.compile(r"^cell ([A-Z]+[0-9]+) \|", re.MULTILINE)


def _diff_file(sheets: dict[str, list[tuple[str, str]]]) -> DiffFile:
    return DiffFile.model_validate(
        {
            "spec_version": 2,
            "task_tolerance": 0.01,
            "sheets": {
                name: {
                    "sheet_weight": 1,
                    "groups": {
                        "intermediate": {
                            "weight": 1,
                            "cells": [{"cell": ref, "cell_type": ctype} for ref, ctype in cells],
                        }
                    },
                }
                for name, cells in sheets.items()
            },
        }
    )


def _bundle_for(tmp_path: Path, complete_path: Path, diff: DiffFile) -> TaskBundle:
    return TaskBundle(
        task_dir=tmp_path,
        mode="predict",
        raw_diff=diff,
        golden=None,
        instructions="Build the model per the brief.",
        init_path=complete_path,
        complete_path=complete_path,
    )


_STATEMENT_CANDS: list[tuple[str, str]] = [
    ("D4", "number"),
    ("E4", "number"),
    ("F4", "number"),
    ("D5", "number"),
    ("E5", "number"),
    ("F5", "number"),
    ("D6", "number"),
    ("D7", "number"),
    ("E7", "number"),
    ("F7", "number"),
    ("D11", "text"),
]


def _statement_setup(tmp_path: Path) -> tuple[TaskBundle, WorkbookData, TaskFeatures]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "MODEL"
    ws["A1"].font = Font(bold=True)
    ws["A3"] = "REVENUE"
    ws["A4"] = "Product A"
    ws["D4"], ws["E4"], ws["F4"] = 100, 110, 121
    ws["A5"] = "Dup"
    ws["D5"], ws["E5"], ws["F5"] = 7, 7, 7
    ws["A6"] = "Precise"
    ws["D6"] = 1.23456789
    ws["A7"] = "Total"
    for col in ("D", "E", "F"):
        ws[f"{col}7"] = f"=SUM({col}4:{col}6)"

    ws["A10"] = "SUMMARY"
    ws["A11"] = "Status"
    ws["D11"] = "DONE"
    ws["H20"] = 4242
    path = tmp_path / "statement.xlsx"
    wb.save(path)
    complete = load_workbook_data(path, cache_dir=tmp_path / "cache")
    bundle = _bundle_for(tmp_path, path, _diff_file({"Model": _STATEMENT_CANDS}))
    return bundle, complete, build_task_features(bundle, complete)


def _dashboard_setup(tmp_path: Path) -> tuple[TaskBundle, WorkbookData, TaskFeatures]:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dash"
    for row in range(2, 10):
        ws[f"B{row}"] = f"Left metric {row}"
        ws[f"C{row}"] = "=A1*2"
        ws[f"E{row}"] = f"Right metric {row}"
        ws[f"F{row}"] = "=SUM(A1:A5)"
    path = tmp_path / "dash.xlsx"
    wb.save(path)
    cands: list[tuple[str, str]] = []
    for row in range(2, 10):
        cands.append((f"C{row}", "number"))
        cands.append((f"F{row}", "number"))
    complete = load_workbook_data(path, cache_dir=tmp_path / "cache")
    bundle = _bundle_for(tmp_path, path, _diff_file({"Dash": cands}))
    return bundle, complete, build_task_features(bundle, complete)


def _build(
    bundle: TaskBundle,
    complete: WorkbookData,
    features: TaskFeatures,
    sheet: str,
    variant: VariantConfig | None = None,
    profile: TaskProfile | None = None,
) -> SheetContext:
    return SheetContext.build(
        bundle=bundle,
        complete=complete,
        features=features,
        sheet_name=sheet,
        variant=variant or VariantConfig(name="test"),
        profile=profile,
    )


def test_choose_grouping_statement_is_row(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    entries = bundle.candidate_cells()["Model"]
    assert choose_grouping(complete.sheets["Model"], entries, features.sheets["Model"]) == "row"


def test_choose_grouping_dashboard_is_cell(tmp_path: Path) -> None:
    bundle, complete, features = _dashboard_setup(tmp_path)
    entries = bundle.candidate_cells()["Dash"]
    assert choose_grouping(complete.sheets["Dash"], entries, features.sheets["Dash"]) == "cell"


def test_block_grouping_is_rejected(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    variant = VariantConfig(name="bad", grouping="block")
    with pytest.raises(ValueError, match="block"):
        _build(bundle, complete, features, "Model", variant)


def test_forced_grouping_overrides_auto(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    ctx = _build(bundle, complete, features, "Model", VariantConfig(name="forced", grouping="cell"))
    assert ctx.grouping == "cell"


def test_context_fences_and_instructions(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    text = _build(bundle, complete, features, "Model").text
    assert "BEGIN SHEET DATA" in text and "END SHEET DATA" in text
    assert "BEGIN TASK INSTRUCTIONS" in text and "END TASK INSTRUCTIONS" in text
    assert "Build the model per the brief." in text
    assert text.index("BEGIN TASK INSTRUCTIONS") < text.index("BEGIN SHEET DATA")


def test_context_each_candidate_row_exactly_once(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    ctx = _build(bundle, complete, features, "Model")
    rows = [int(m) for m in ROW_LINE_RE.findall(ctx.text)]
    assert sorted(rows) == [4, 5, 6, 7, 11]
    assert len(rows) == len(set(rows))
    assert sorted(ctx.candidate_refs) == sorted(e.cell for e in bundle.candidate_cells()["Model"])


def test_context_compression_rules(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    text = _build(bundle, complete, features, "Model").text
    assert "D5:F5=7" in text
    assert "1.2346" in text and "1.23456789" not in text
    assert "Product A" in text
    assert "4242" not in text
    assert re.search(r"^row 7 \| Total \| =SUM\(R\[-3\]C:R\[-1\]C\)", text, re.MULTILINE)


def test_context_workbook_map_and_sections(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    text = _build(bundle, complete, features, "Model").text
    map_start = text.index("WORKBOOK MAP")
    assert text.index("BEGIN SHEET DATA") < map_start
    assert re.search(r"(?m)^- Model \| used A1:H20 .* candidates 11", text)
    assert "rows 3-7: REVENUE" in text
    assert "rows 10-11: SUMMARY" in text


def test_context_profile_lines_only_when_provided(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    profile = TaskProfile(
        workbook_purpose="LBO of a target",
        model_type="LBO",
        expected_deliverables=["returns schedule"],
        likely_capstone_outputs=["IRR", "MOIC"],
    )
    with_profile = _build(bundle, complete, features, "Model", profile=profile).text
    without = _build(bundle, complete, features, "Model").text
    assert "LBO of a target" in with_profile and "MOIC" in with_profile
    assert "TASK PROFILE" in with_profile
    assert "TASK PROFILE" not in without


def test_context_feature_hints_flagged_by_variant(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    plain = _build(bundle, complete, features, "Model").text
    hinted = _build(bundle, complete, features, "Model", VariantConfig(name="v2", features_in_prompt=True)).text
    assert "deps_out=" in hinted and "agg=SUM" in hinted
    assert "deps_out=" not in plain


def test_text_candidate_row_encoded(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    text = _build(bundle, complete, features, "Model").text
    assert re.search(r"(?m)^row 11 \| Status \| - \| D11=DONE \| text", text)


def test_cell_mode_lines_each_candidate_once(tmp_path: Path) -> None:
    bundle, complete, features = _dashboard_setup(tmp_path)
    ctx = _build(bundle, complete, features, "Dash")
    assert ctx.grouping == "cell"
    refs = CELL_LINE_RE.findall(ctx.text)
    assert sorted(refs) == sorted(e.cell for e in bundle.candidate_cells()["Dash"])
    assert len(refs) == len(set(refs))
    assert "Left metric 2" in ctx.text


def test_split_by_sections_partitions_on_boundaries(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    ctx = _build(bundle, complete, features, "Model")
    parts = split_by_sections(ctx, threshold_rows=4)
    assert len(parts) == 2
    assert [p.part for p in parts] == [1, 2] and all(p.parts == 2 for p in parts)
    assert sorted({cell_row(r) for r in parts[0].candidate_refs}) == [4, 5, 6, 7]
    assert sorted({cell_row(r) for r in parts[1].candidate_refs}) == [11]
    combined = [ref for p in parts for ref in p.candidate_refs]
    assert sorted(combined) == sorted(ctx.candidate_refs)
    assert len(combined) == len(set(combined))
    assert "PART 1/2" in parts[0].text and "PART 2/2" in parts[1].text


def test_split_below_threshold_returns_context_unchanged(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    ctx = _build(bundle, complete, features, "Model")
    assert split_by_sections(ctx, threshold_rows=50) == [ctx]


def test_split_oversized_single_section_stays_whole(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    ctx = _build(bundle, complete, features, "Model")
    parts = split_by_sections(ctx, threshold_rows=1)
    revenue_part = parts[0]
    assert sorted({cell_row(r) for r in revenue_part.candidate_refs}) == [4, 5, 6, 7]


def test_full_dump_bypasses_compression(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    ctx = _build(bundle, complete, features, "Model", VariantConfig(name="v5", full_dump=True))
    assert "4242" in ctx.text
    assert sorted(ctx.candidate_refs) == sorted(e.cell for e in bundle.candidate_cells()["Model"])


def test_token_estimate_is_chars_div_4(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    ctx = _build(bundle, complete, features, "Model")
    assert estimate_tokens("x" * 41) == 10
    assert ctx.token_estimate == len(ctx.text) // 4


def test_build_requires_candidates(tmp_path: Path) -> None:
    bundle, complete, features = _statement_setup(tmp_path)
    with pytest.raises(ValueError, match="candidate"):
        _build(bundle, complete, features, "Sheet-with-no-candidates")
