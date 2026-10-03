from __future__ import annotations

import random
import re
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]
import pytest
from openpyxl.styles import Color, Font  # type: ignore[import-untyped]
from openpyxl.workbook.defined_name import DefinedName  # type: ignore[import-untyped]

from formuloom.bundle import TaskBundle
from formuloom.features import (
    DOMAIN_LEXICON,
    RowFeatures,
    Section,
    SheetNameMismatchError,
    build_reference_graph,
    build_task_features,
    detect_sections,
    formula_sketch,
    to_r1c1,
)
from formuloom.schema import DiffFile
from formuloom.workbook import load_workbook_data

DATA_ROOT = Path(__file__).resolve().parents[1] / "data"
TASKS = ["synthetic-statement", "synthetic-budget", "synthetic-rollup", "synthetic-projection", "synthetic-summary"]

FEATURES_ALL_TASKS_BUDGET_SECONDS = 60.0


def _diff_file(sheets: dict[str, list[tuple[str, str]]]) -> DiffFile:
    return DiffFile.model_validate(
        {
            "spec_version": 2,
            "task_tolerance": 0.01,
            "sheets": {
                name: {
                    "sheet_weight": 1,
                    "groups": {
                        "intermediate": {
                            "weight": 1,
                            "cells": [{"cell": ref, "cell_type": ctype} for ref, ctype in cells],
                        }
                    },
                }
                for name, cells in sheets.items()
            },
        }
    )


def _bundle_for(tmp_path: Path, complete_path: Path, diff: DiffFile) -> TaskBundle:
    return TaskBundle(
        task_dir=tmp_path,
        mode="predict",
        raw_diff=diff,
        golden=None,
        instructions="Build the model per the brief.",
        init_path=complete_path,
        complete_path=complete_path,
    )


def _graph_workbook(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    data = wb.active
    data.title = "Data"
    data["A1"] = 10
    data["A2"] = 20
    data["B1"], data["B2"], data["B3"] = 1, 2, 3
    for r in range(1, 8):
        data[f"C{r}"] = r
    data["E1"] = "='Calc M'!A5*2"

    calc = wb.create_sheet("Calc M")
    calc["A5"] = 5
    calc["B2"] = "=Data!A1+A5"
    calc["C2"] = "=SUM(Data!B1:B3)"
    calc["D2"] = "=MyTotal*2"
    calc["E2"] = "=SUM(Data!C:C)"
    calc["F2"] = "=[Book1]Sheet1!A1+Data!A2"
    calc["G2"] = "=Table1[Col]"
    calc["H2"] = "=)("
    wb.defined_names.add(DefinedName("MyTotal", attr_text="Data!$B$1:$B$3"))
    path = tmp_path / "graph.xlsx"
    wb.save(path)
    return path


def _statement_workbook(tmp_path: Path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Model"
    ws["A1"] = "INCOME STATEMENT"
    ws["A1"].font = Font(bold=True)

    ws["A3"] = "REVENUE BUILD"
    ws["A4"] = "Product A"
    ws["D4"], ws["E4"], ws["F4"] = 100, 110, 121
    ws["D4"].font = Font(color=Color(rgb="FF0000FF"))
    ws["A5"] = "Product B"
    for col in ("D", "E", "F"):
        ws[f"{col}5"] = f"={col}4*2"
    ws["A6"] = "Total Revenue"
    ws["A6"].font = Font(bold=True)
    for col in ("D", "E", "F"):
        ws[f"{col}6"] = f"=SUM({col}4:{col}5)"

    ws["A9"] = "Summary"
    ws["A9"].font = Font(bold=True)
    ws["A10"] = "Net Income"
    ws["D10"] = "=D6*0.5+Aux!B1"
    ws["A11"] = "Status"
    ws["D11"] = "DONE"

    aux = wb.create_sheet("Aux")
    aux["B1"] = 99
    aux["C1"] = "=Model!D6"
    wb.defined_names.add(DefinedName("NetIncomeCell", attr_text="Model!$D$10"))
    path = tmp_path / "statement.xlsx"
    wb.save(path)
    return path


_STATEMENT_CANDIDATES: list[tuple[str, str]] = [
    ("D4", "number"),
    ("E4", "number"),
    ("F4", "number"),
    ("D5", "number"),
    ("E5", "number"),
    ("F5", "number"),
    ("D6", "number"),
    ("E6", "number"),
    ("F6", "number"),
    ("D10", "number"),
    ("D11", "text"),
]


def _statement_features(tmp_path: Path) -> dict[int, RowFeatures]:
    path = _statement_workbook(tmp_path)
    complete = load_workbook_data(path, cache_dir=tmp_path / "cache")
    bundle = _bundle_for(tmp_path, path, _diff_file({"Model": _STATEMENT_CANDIDATES}))
    feats = build_task_features(bundle, complete)
    return {r.row: r for r in feats.sheets["Model"].rows}


def test_precedents_and_dependents_by_name(tmp_path: Path) -> None:
    data = load_workbook_data(_graph_workbook(tmp_path), cache_dir=tmp_path / "cache")
    graph = build_reference_graph(data)

    assert graph.precedents[("Calc M", "B2")] == {("Data", "A1"), ("Calc M", "A5")}
    assert graph.precedents[("Calc M", "C2")] == {("Data", "B1"), ("Data", "B2"), ("Data", "B3")}
    assert graph.precedents[("Data", "E1")] == {("Calc M", "A5")}

    assert graph.dependents[("Data", "A1")] == {("Calc M", "B2")}
    assert ("Calc M", "B2") in graph.dependents[("Calc M", "A5")]
    assert ("Data", "E1") in graph.dependents[("Calc M", "A5")]

    assert graph.precedents.get(("Data", "A1"), set()) == set()
    assert graph.dependents.get(("Calc M", "B2"), set()) == set()


def test_defined_name_resolves_to_target_range(tmp_path: Path) -> None:
    data = load_workbook_data(_graph_workbook(tmp_path), cache_dir=tmp_path / "cache")
    graph = build_reference_graph(data)
    assert graph.precedents[("Calc M", "D2")] == {("Data", "B1"), ("Data", "B2"), ("Data", "B3")}


def test_whole_column_ref_capped_with_warning(tmp_path: Path) -> None:
    data = load_workbook_data(_graph_workbook(tmp_path), cache_dir=tmp_path / "cache")
    graph = build_reference_graph(data)
    expected = {("Data", f"C{r}") for r in range(1, 8)}
    assert graph.precedents[("Calc M", "E2")] == expected
    assert any("capped" in w for w in graph.warnings[("Calc M", "E2")])


def test_unresolvable_refs_warn_never_crash(tmp_path: Path) -> None:
    data = load_workbook_data(_graph_workbook(tmp_path), cache_dir=tmp_path / "cache")
    graph = build_reference_graph(data)

    assert graph.precedents[("Calc M", "F2")] == {("Data", "A2")}
    assert graph.warnings[("Calc M", "F2")]

    assert graph.precedents[("Calc M", "G2")] == set()
    assert graph.warnings[("Calc M", "G2")]

    assert graph.precedents[("Calc M", "H2")] == set()
    assert graph.warnings[("Calc M", "H2")]


@pytest.mark.parametrize(
    ("formula", "row", "col", "expected"),
    [
        ("=A1", 2, 2, "=R[-1]C[-1]"),
        ("=A1", 1, 1, "=RC"),
        ("=$A$1", 7, 9, "=R1C1"),
        ("=$A1", 3, 4, "=R[-2]C1"),
        ("=A$1", 3, 4, "=R1C[-3]"),
        ("=SUM(F9:F20)", 21, 6, "=SUM(R[-12]C:R[-1]C)"),
        ("='My Sheet'!A1", 1, 2, "='My Sheet'!RC[-1]"),
        ("=Sheet2!$C$3", 1, 1, "=Sheet2!R3C3"),
        ("=MyName*2", 5, 5, "=MyName*2"),
        ("=SUM(F:F)", 4, 3, "=SUM(C[3])"),
        ("=SUM($3:$3)", 9, 9, "=SUM(R3)"),
    ],
)
def test_to_r1c1_cases(formula: str, row: int, col: int, expected: str) -> None:
    assert to_r1c1(formula, row, col) == expected


@pytest.mark.parametrize(
    "formula",
    ["=Table1[Col]", "=[Book1]Sheet1!A1", "=)(", "=@#$%"],
)
def test_to_r1c1_unparseable_returns_none(formula: str) -> None:
    assert to_r1c1(formula, 5, 5) is None


_R1C1_CELL_RE = re.compile(r"^R(?:(\d+)|\[(-?\d+)\])?C(?:(\d+)|\[(-?\d+)\])?$")


def _decode_r1c1(part: str, origin_row: int, origin_col: int) -> tuple[int, int]:
    match = _R1C1_CELL_RE.match(part)
    assert match is not None, part
    r_abs, r_rel, c_abs, c_rel = match.groups()
    row = int(r_abs) if r_abs else origin_row + int(r_rel) if r_rel else origin_row
    col = int(c_abs) if c_abs else origin_col + int(c_rel) if c_rel else origin_col
    return row, col


def _col_letters(idx: int) -> str:
    out = ""
    while idx:
        idx, rem = divmod(idx - 1, 26)
        out = chr(ord("A") + rem) + out
    return out


def test_to_r1c1_roundtrip_property() -> None:
    rng = random.Random(20260707)
    for _ in range(300):
        o_row, o_col = rng.randint(1, 400), rng.randint(1, 150)
        t_row, t_col = rng.randint(1, 400), rng.randint(1, 150)
        r_abs, c_abs = rng.random() < 0.5, rng.random() < 0.5
        sheet = rng.choice(["", "Sheet2!", "'My Sheet'!"])
        a1 = f"{'$' if c_abs else ''}{_col_letters(t_col)}{'$' if r_abs else ''}{t_row}"
        result = to_r1c1(f"={sheet}{a1}", o_row, o_col)
        assert result is not None, (sheet, a1, o_row, o_col)
        body = result[1:]
        assert body.startswith(sheet)
        decoded = _decode_r1c1(body[len(sheet) :], o_row, o_col)
        assert decoded == (t_row, t_col), (sheet, a1, o_row, o_col, result)


def test_to_r1c1_translation_invariance() -> None:
    rng = random.Random(7)
    for _ in range(100):
        d_row, d_col = rng.randint(-5, 5), rng.randint(-5, 5)
        row_a, col_a = rng.randint(10, 50), rng.randint(10, 30)
        row_b, col_b = rng.randint(60, 100), rng.randint(40, 60)
        ref_a = f"{_col_letters(col_a + d_col)}{row_a + d_row}"
        ref_b = f"{_col_letters(col_b + d_col)}{row_b + d_row}"
        assert to_r1c1(f"={ref_a}*2", row_a, col_a) == to_r1c1(f"={ref_b}*2", row_b, col_b)


@pytest.mark.parametrize(
    ("formula", "expected"),
    [
        ("=SUM(F9:F20)", "SUM(REF:REF)"),
        ('=IF(A1>0,"yes",1.5)', "IF(REF>NUM,STR,NUM)"),
        ("=VLOOKUP($F4,'Fin St'!$B$4:$N$40,7,FALSE)", "VLOOKUP(REF,REF:REF,NUM,FALSE)"),
        ("=MyName*2", "REF*NUM"),
        ("=D6*0.5+Aux!B1", "REF*NUM+REF"),
    ],
)
def test_formula_sketch_cases(formula: str, expected: str) -> None:
    assert formula_sketch(formula) == expected


def test_formula_sketch_malformed_still_returns() -> None:
    for bad in ("=)(", "=@#$%", "=1+", "="):
        out = formula_sketch(bad)
        assert isinstance(out, str)
    assert "NUM" in formula_sketch("=1+")


def test_detect_sections_statement_sheet(tmp_path: Path) -> None:
    data = load_workbook_data(_statement_workbook(tmp_path), cache_dir=tmp_path / "cache")
    sections = detect_sections(data.sheets["Model"])
    assert [(s.title, s.start_row, s.end_row) for s in sections] == [
        ("INCOME STATEMENT", 1, 1),
        ("REVENUE BUILD", 3, 6),
        ("Summary", 9, 11),
    ]
    assert all(isinstance(s, Section) for s in sections)


def test_detect_sections_untitled_lead_and_blank_termination(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "S"
    ws["A1"] = "just data"
    ws["B1"] = 5
    ws["A2"] = "more"
    ws["B2"] = 6

    ws["A5"] = "TOTALS"
    ws["A5"].font = Font(bold=True)
    ws["A6"] = "x"
    ws["B6"] = 1
    path = tmp_path / "sections.xlsx"
    wb.save(path)
    data = load_workbook_data(path, cache_dir=tmp_path / "cache")
    sections = detect_sections(data.sheets["S"])
    assert [(s.title, s.start_row, s.end_row) for s in sections] == [
        (None, 1, 2),
        ("TOTALS", 5, 6),
    ]


def test_detect_sections_empty_sheet(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "Empty"
    wb.create_sheet("Full")["A1"] = "x"
    path = tmp_path / "empty.xlsx"
    wb.save(path)
    data = load_workbook_data(path, cache_dir=tmp_path / "cache")
    assert detect_sections(data.sheets["Empty"]) == []


def test_row_features_labels_sections_and_counts(tmp_path: Path) -> None:
    rows = _statement_features(tmp_path)
    assert set(rows) == {4, 5, 6, 10, 11}
    assert rows[4].label == "Product A"
    assert rows[4].section_title == "REVENUE BUILD"
    assert rows[4].n_candidates == 3
    assert rows[10].section_title == "Summary"


def test_row_features_input_color_and_style(tmp_path: Path) -> None:
    rows = _statement_features(tmp_path)
    assert rows[4].input_colored_fraction == pytest.approx(1 / 3)
    assert rows[5].input_colored_fraction == 0.0
    assert rows[6].bold_or_bordered is False
    assert rows[4].bold_or_bordered is False


def test_row_features_patterns(tmp_path: Path) -> None:
    rows = _statement_features(tmp_path)
    assert rows[4].pattern_kind == "none"
    assert rows[4].formula_pattern is None
    assert rows[5].pattern_kind == "r1c1"
    assert rows[5].formula_pattern == "=R[-1]C*2"
    assert rows[6].pattern_kind == "r1c1"
    assert rows[6].formula_pattern == "=SUM(R[-2]C:R[-1]C)"


def test_row_features_aggregation_and_lexicon(tmp_path: Path) -> None:
    rows = _statement_features(tmp_path)
    assert rows[6].aggregates_range is True
    assert rows[6].agg_function == "SUM"
    assert "total" in rows[6].lexicon_hits
    assert rows[5].aggregates_range is False
    assert rows[5].agg_function is None
    assert "net income" in rows[10].lexicon_hits


def test_row_features_graph_derived(tmp_path: Path) -> None:
    rows = _statement_features(tmp_path)

    assert rows[6].n_dependents_outside_row == 2

    assert rows[10].distinct_source_sheets == 2
    assert rows[10].distinct_source_sections == 2
    assert rows[6].distinct_source_sheets == 1
    assert rows[6].distinct_source_sections == 1


def test_row_features_presentation_flags(tmp_path: Path) -> None:
    rows = _statement_features(tmp_path)
    assert rows[10].in_named_range is True
    assert rows[5].in_named_range is False


def test_text_candidate_row_gets_full_features(tmp_path: Path) -> None:
    rows = _statement_features(tmp_path)
    row = rows[11]
    assert row.label == "Status"
    assert row.section_title == "Summary"
    assert row.pattern_kind == "none"
    assert row.n_candidates == 1
    assert row.candidate_refs == ["D11"]


def test_malformed_formula_row_warns_not_crashes(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "M"
    ws["A1"] = "Broken"
    ws["B1"] = "=)("
    path = tmp_path / "bad.xlsx"
    wb.save(path)
    complete = load_workbook_data(path, cache_dir=tmp_path / "cache")
    bundle = _bundle_for(tmp_path, path, _diff_file({"M": [("B1", "number")]}))
    feats = build_task_features(bundle, complete)
    row = feats.sheets["M"].rows[0]
    assert row.pattern_kind in ("sketch", "none", "mixed")
    assert row.warnings


def test_sheet_name_mismatch_is_loud_typed_error(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active.title = "Revenue"
    wb.active["A1"] = 1
    path = tmp_path / "mm.xlsx"
    wb.save(path)
    complete = load_workbook_data(path, cache_dir=tmp_path / "cache")
    bundle = _bundle_for(tmp_path, path, _diff_file({"revenue ": [("A1", "number")]}))
    with pytest.raises(SheetNameMismatchError) as exc:
        build_task_features(bundle, complete)
    assert "Revenue" in str(exc.value)


def test_lexicon_is_domain_derived() -> None:
    for term in ("total", "ebitda", "irr", "moic", "dscr", "enterprise value", "ending cash"):
        assert term in DOMAIN_LEXICON
    for leaked in ("synthetic-person-id", "fixture-only-token", "unapproved-name"):
        assert all(leaked not in term for term in DOMAIN_LEXICON)
