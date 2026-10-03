import asyncio
import json
from pathlib import Path

import openpyxl
import pytest
from typer.testing import CliRunner

from formuloom.bundle import PREDICT_ALLOWLIST, LeakageError, TaskBundle, _read_text
from formuloom.cli import app
from formuloom.fixtures import generate_fixture
from formuloom.offline import inspect_bundle, run_mechanical_v15
from formuloom.weak import predict_v8


def test_synthetic_pair_has_real_cross_sheet_formulas_and_cached_values(tmp_path):
    root = generate_fixture(tmp_path / "example")
    formula = openpyxl.load_workbook(root / "complete.xlsx", data_only=False)
    values = openpyxl.load_workbook(root / "complete.xlsx", data_only=True)
    assert formula["Summary"]["B2"].value == "=Operations!B5"
    assert values["Summary"]["B2"].value == 660
    assert values["Operations"]["B5"].value == 660
    assert formula.properties.creator == "Formuloom"
    assert len(formula.sheetnames) == 3
    formula.close()
    values.close()


def test_fixture_refuses_overwrite(tmp_path):
    root = generate_fixture(tmp_path / "example")
    with pytest.raises(FileExistsError):
        generate_fixture(root)


def test_weak_prediction_never_reads_reference_labels(tmp_path, monkeypatch):
    root = generate_fixture(tmp_path / "example")
    accessed = []
    original = Path.open

    def spy(path, *args, **kwargs):
        accessed.append(path.name)
        return original(path, *args, **kwargs)

    monkeypatch.setattr(Path, "open", spy)
    bundle = TaskBundle.load(root, "predict")
    diff, diagnostics = predict_v8(bundle, cache_dir=tmp_path / "cache")
    assert "raw_diff.json" in accessed
    assert "subset.json" not in accessed
    assert diagnostics["lf_names"]
    assert diagnostics["em"]["n_rows"] > 0
    assert all(0 <= p <= 1 for p in diagnostics["probabilities"])
    assert set(diff.sheets) == set(bundle.raw_diff.sheets)
    with pytest.raises(LeakageError):
        _read_text(root / "subset.json", PREDICT_ALLOWLIST)


def test_weak_threshold_is_meaningful_and_validated(tmp_path):
    bundle = TaskBundle.load(generate_fixture(tmp_path / "example"), "predict")
    low, _ = predict_v8(bundle, threshold=0, cache_dir=tmp_path / "cache")
    high, _ = predict_v8(bundle, threshold=1, cache_dir=tmp_path / "cache")
    assert sum(len(low.final_refs(s)) for s in low.sheets) > 0
    assert not any(high.final_refs(s) for s in high.sheets)
    for value in (-0.1, 1.1, float("nan")):
        with pytest.raises(ValueError):
            predict_v8(bundle, threshold=value)


def test_weak_output_and_diagnostics_are_repeatable(tmp_path):
    bundle = TaskBundle.load(generate_fixture(tmp_path / "example"), "predict")
    assert predict_v8(bundle, cache_dir=tmp_path / "cache") == predict_v8(bundle, cache_dir=tmp_path / "cache")


def test_mechanical_v15_runs_real_composition_and_pruning_without_provider(tmp_path, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("Provider construction is forbidden")

    monkeypatch.setattr("formuloom.cli.build_async_client", forbidden)
    bundle = TaskBundle.load(generate_fixture(tmp_path / "example"), "predict")
    result, routes = asyncio.run(run_mechanical_v15(bundle, run_dir=tmp_path / "run"))
    assert result.usage.api_calls == 0
    assert result.usage.cost_usd == 0
    assert set(routes) == set(bundle.raw_diff.sheets)
    for sheet in result.diff.sheets:
        universe = {c.cell for c in bundle.candidate_cells()[sheet]}
        assert result.diff.final_refs(sheet) | result.diff.intermediate_refs(sheet) == universe
        assert not result.diff.final_refs(sheet) & result.diff.intermediate_refs(sheet)
    assert routes["Assumptions"]["emptied"] is False


def test_inspection_contains_actual_graph_and_em_diagnostics(tmp_path):
    bundle = TaskBundle.load(generate_fixture(tmp_path / "example"), "predict")
    report = inspect_bundle(bundle, cache_dir=tmp_path / "cache")
    assert report["schema_version"] == 1
    assert report["synthetic"] is True
    assert report["sheets"]["Operations"]["cross_sheet_dependents"] > 0
    assert report["weak_supervision"]["em"]["n_lfs"] == 7


def test_public_cli_weak_and_offline_evaluation(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    root = generate_fixture(tmp_path / "data" / "example")
    runner = CliRunner()
    result = runner.invoke(app, ["predict", str(root), "--variant", "V8", "-o", str(tmp_path / "v8.json")])
    assert result.exit_code == 0, result.output
    result = runner.invoke(
        app, ["eval", str(root.parent), "--variant", "V8", "--offline-only", "--out", str(tmp_path / "scores.json")]
    )
    assert result.exit_code == 0, result.output
    payload = json.loads((tmp_path / "scores.json").read_text())
    assert payload["meta"]["tokens"] == 0
    assert payload["meta"]["offline_only"] is True
    assert payload["meta"]["variant"] == "V8"


def test_reproduction_runs_actual_predictions_and_refuses_overwrite(tmp_path, monkeypatch):
    from formuloom.reproduce import reproduce

    def forbidden(*args, **kwargs):
        raise AssertionError("Provider construction is forbidden")

    monkeypatch.setattr("formuloom.cli.build_async_client", forbidden)
    report = reproduce(tmp_path / "repro")
    assert report["synthetic"] is True
    assert report["provider_calls"] == 0
    assert set(report["predictions"]) == {"V0", "V8", "V15-mechanical"}
    for name, prediction in report["predictions"].items():
        from formuloom.schema import DiffFile
        from formuloom.score import score_task

        predicted = DiffFile.model_validate_json((tmp_path / "repro" / f"{name}.json").read_text())
        reference = DiffFile.model_validate_json((tmp_path / "repro" / "fixture" / "subset.json").read_text())
        assert prediction["strict"] == score_task("synthetic", predicted, reference, "strict").model_dump()
    with pytest.raises(FileExistsError):
        reproduce(tmp_path / "repro")
