from __future__ import annotations

import json
from pathlib import Path

import openpyxl  # type: ignore[import-untyped]
from openpyxl.styles import Color, Font  # type: ignore[import-untyped]

from formuloom.workbook import (
    CellRecord,
    SheetData,
    WorkbookData,
    WorkbookPair,
    apply_tint,
    get_parse_count,
    is_blue_input,
    load_workbook_data,
    parse_theme_colors,
    reset_parse_count,
    resolve_font_rgb,
    value_diff,
)

DATA_ROOT = Path(__file__).resolve().parents[1] / "data"
TASKS = ["synthetic-statement", "synthetic-budget", "synthetic-rollup", "synthetic-projection", "synthetic-summary"]


def _save(wb: openpyxl.Workbook, path: Path) -> Path:
    wb.save(path)
    return path


def test_apply_tint_known_values() -> None:
    assert apply_tint("000000", 0.5) == "808080"
    assert apply_tint("FFFFFF", -0.25) == "BFBFBF"
    assert apply_tint("1F497D", 0.0) == "1F497D"


def test_parse_theme_colors_applies_index_swap() -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    import io

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    reloaded = openpyxl.load_workbook(buf)
    theme = parse_theme_colors(reloaded.loaded_theme)
    assert theme[0] == "FFFFFF"
    assert theme[1] == "000000"
    assert theme[3] == "1F497D"
    assert theme[4] == "4F81BD"


def test_parse_theme_colors_missing_returns_empty() -> None:
    assert parse_theme_colors(None) == []
    assert parse_theme_colors(b"<no scheme here/>") == []


def test_resolve_font_rgb_all_three_encodings() -> None:
    theme_colors = ["FFFFFF", "000000", "EEECE1", "1F497D", "4F81BD"]

    assert resolve_font_rgb(Color(rgb="FF0000FF"), theme_colors) == "0000FF"

    assert resolve_font_rgb(Color(indexed=4), theme_colors) == "0000FF"

    assert resolve_font_rgb(Color(theme=3, tint=0.0), theme_colors) == "1F497D"

    tinted = resolve_font_rgb(Color(theme=3, tint=0.4), theme_colors)
    assert tinted is not None and is_blue_input(tinted)


def test_resolve_font_rgb_unresolvable_is_none() -> None:
    assert resolve_font_rgb(None, []) is None

    assert resolve_font_rgb(Color(theme=4, tint=0.0), []) is None


def test_is_blue_input_heuristic() -> None:
    assert is_blue_input("0000FF") is True
    assert is_blue_input("1F4E79") is True
    assert is_blue_input("000000") is False
    assert is_blue_input("FF0000") is False
    assert is_blue_input("808080") is False
    assert is_blue_input(None) is False
    assert is_blue_input("bad") is False


def test_color_resolution_end_to_end(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "rgb"
    ws["A1"].font = Font(color=Color(rgb="FF0000FF"))
    ws["A2"] = "theme"
    ws["A2"].font = Font(color=Color(theme=3, tint=0.4))
    ws["A3"] = "indexed"
    ws["A3"].font = Font(color=Color(indexed=4))
    ws["A4"] = "black"
    ws["A4"].font = Font(color=Color(rgb="FF000000"))
    path = _save(wb, tmp_path / "colors.xlsx")

    data = load_workbook_data(path, cache_dir=tmp_path / "cache")
    cells = data.sheets["Sheet"].cells
    assert cells["A1"].font_rgb == "0000FF" and cells["A1"].is_input_colored
    assert cells["A2"].is_input_colored
    assert cells["A3"].font_rgb == "0000FF" and cells["A3"].is_input_colored
    assert cells["A4"].font_rgb == "000000" and not cells["A4"].is_input_colored


def test_dual_view_formula_and_value(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=B1+C1"
    ws["B1"] = 5
    ws["C1"] = 7
    path = _save(wb, tmp_path / "formula.xlsx")

    data = load_workbook_data(path, cache_dir=tmp_path / "cache")
    cells = data.sheets["Sheet"].cells
    assert cells["A1"].formula == "=B1+C1"
    assert cells["A1"].value is None
    assert cells["B1"].formula is None and cells["B1"].value == 5


def test_merged_cells_anchor(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.merge_cells("A1:C1")
    ws["A1"] = "Header"
    ws["A2"] = "body"
    path = _save(wb, tmp_path / "merged.xlsx")

    cells = load_workbook_data(path, cache_dir=tmp_path / "cache").sheets["Sheet"].cells
    assert cells["A1"].merged_anchor == "A1"
    assert cells["A2"].merged_anchor is None

    assert "B1" not in cells and "C1" not in cells


def test_bold_and_borders(tmp_path: Path) -> None:
    from openpyxl.styles import Border, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "total"
    ws["A1"].font = Font(bold=True)
    ws["A1"].border = Border(top=Side(style="thin"), bottom=Side(style="double"))
    ws["A2"] = "plain"
    path = _save(wb, tmp_path / "style.xlsx")

    cells = load_workbook_data(path, cache_dir=tmp_path / "cache").sheets["Sheet"].cells
    assert cells["A1"].bold and cells["A1"].border_top and cells["A1"].border_bottom
    assert not cells["A2"].bold and not cells["A2"].border_top


def test_hidden_rows_cols_freeze_and_sheet_state(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Visible"
    ws["A1"] = "x"
    ws.freeze_panes = "B2"
    ws.row_dimensions[2].hidden = True
    ws.column_dimensions["B"].hidden = True
    hidden_ws = wb.create_sheet("Secret")
    hidden_ws["A1"] = "y"
    hidden_ws.sheet_state = "hidden"
    path = _save(wb, tmp_path / "layout.xlsx")

    data = load_workbook_data(path, cache_dir=tmp_path / "cache")
    visible = data.sheets["Visible"]
    assert visible.frozen_panes == "B2"
    assert visible.hidden_rows == {2}
    assert visible.hidden_cols == {"B"}
    assert visible.hidden is False
    assert data.sheets["Secret"].hidden is True


def test_defined_names_captured(tmp_path: Path) -> None:
    from openpyxl.workbook.defined_name import DefinedName  # type: ignore[import-untyped]

    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.defined_names.add(DefinedName("MyName", attr_text="Sheet!$A$1"))
    path = _save(wb, tmp_path / "names.xlsx")

    data = load_workbook_data(path, cache_dir=tmp_path / "cache")
    assert data.defined_names.get("MyName") == "Sheet!$A$1"


def test_dataclass_shapes_are_stable() -> None:
    assert set(CellRecord.__dataclass_fields__) == {
        "sheet",
        "ref",
        "row",
        "col",
        "value",
        "formula",
        "font_rgb",
        "is_input_colored",
        "bold",
        "border_top",
        "border_bottom",
        "number_format",
        "merged_anchor",
    }
    assert set(SheetData.__dataclass_fields__) == {
        "name",
        "cells",
        "max_row",
        "max_col",
        "hidden",
        "frozen_panes",
        "hidden_rows",
        "hidden_cols",
    }
    assert set(WorkbookData.__dataclass_fields__) == {
        "path",
        "sha256",
        "sheets",
        "defined_names",
        "chart_source_refs",
        "print_areas",
        "table_ranges",
    }


def test_value_diff_synthetic(tmp_path: Path) -> None:
    init = openpyxl.Workbook()
    iws = init.active
    iws.title = "S1"
    iws["A1"] = 1
    iws["B1"] = "hello"
    iws["C1"] = 3
    ipath = _save(init, tmp_path / "init.xlsx")

    comp = openpyxl.Workbook()
    cws = comp.active
    cws.title = "S1"
    cws["A1"] = 2
    cws["B1"] = "hello"
    cws["D1"] = 99

    new = comp.create_sheet("S2")
    new["A1"] = 7
    cpath = _save(comp, tmp_path / "complete.xlsx")

    pair = WorkbookPair.load(ipath, cpath, cache_dir=tmp_path / "cache")
    diff = value_diff(pair)
    assert diff["S1"] == {"A1", "C1", "D1"}
    assert diff["S2"] == {"A1"}


def test_cache_hit_avoids_reparse(tmp_path: Path) -> None:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    path = _save(wb, tmp_path / "cacheme.xlsx")
    cache_dir = tmp_path / "cache"

    reset_parse_count()
    first = load_workbook_data(path, cache_dir=cache_dir)
    assert get_parse_count() == 1
    second = load_workbook_data(path, cache_dir=cache_dir)
    assert get_parse_count() == 1
    assert first.sha256 == second.sha256
    assert set(first.sheets) == set(second.sheets)


def test_cache_version_invalidates(tmp_path: Path) -> None:
    path = tmp_path / "v.xlsx"
    cache_dir = tmp_path / "cache"
    wb = openpyxl.Workbook()
    wb.active["A1"] = "one"
    _save(wb, path)
    reset_parse_count()
    load_workbook_data(path, cache_dir=cache_dir)
    wb2 = openpyxl.Workbook()
    wb2.active["A1"] = "two"
    _save(wb2, path)
    load_workbook_data(path, cache_dir=cache_dir)
    assert get_parse_count() == 2


def _raw_diff_union(task: str) -> dict[str, set[str]]:
    raw = json.loads((DATA_ROOT / task / "raw_diff.json").read_text(encoding="utf-8"))
    out: dict[str, set[str]] = {}
    for sheet_name, sheet in raw["sheets"].items():
        refs: set[str] = set()
        for group in sheet["groups"].values():
            for cell in group["cells"]:
                refs.add(cell["cell"])
        out[sheet_name] = refs
    return out
